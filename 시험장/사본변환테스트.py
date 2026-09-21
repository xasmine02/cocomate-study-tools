# -*- coding: utf-8 -*-
"""시험장.py v2.2.1 풀이 사본(xlsx→xlsm) 변환 + Excel 실행 헬퍼 헤드리스 테스트.

  1. 실제 문제 파일(.xlsx, 있을 때) 변환: 모든 항목 flag_bits에 0x08 없음,
     [Content_Types].xml 첫 항목, testzip 통과, 매크로 content-type, 나머지
     파트 바이트 동일, openpyxl 로드
  2. data descriptor(0x08) 비트·CT 순서가 어긋난 xlsx를 만들어 변환하면 깨끗해짐
  3. copy_as_macro_enabled: .xlsm 사본 + 시작 로그 기록, 변환 불가 파일은
     .xlsx 폴백, .xlsm 원본은 그대로 복사
  4. Excel 실행 헬퍼: 명령 문자열 파싱, 잠금 파일 감지, 열림 판정,
     Windows 외에서는 탐색 결과 None/unknown

실제 파일: 환경변수 COCO_TEST_XLSX(문제 .xlsx 경로). 없으면 openpyxl로 만든
합성 xlsx만 사용합니다. 실제 문제 파일은 절대 저장소에 복사하지 않습니다.
"""
import importlib.util
import io
import os
import shutil
import sys
import tempfile
import zipfile

BASE = os.path.dirname(os.path.abspath(__file__))
spec = importlib.util.spec_from_file_location("sj", os.path.join(BASE, "시험장.py"))
sj = importlib.util.module_from_spec(spec)
spec.loader.exec_module(sj)
N = 0
TMP = tempfile.mkdtemp(prefix="coco_conv_")
sj.STARTUP_LOG_PATH = os.path.join(TMP, "시작로그.txt")   # 실제 로그 오염 방지


def check(desc, cond, extra=""):
    global N
    N += 1
    print(f"  [{N:02d}] {desc}: {'OK' if cond else 'FAIL'}" + (f" ({extra})" if extra else ""))
    if not cond:
        raise AssertionError(f"{desc} {extra}")


def real_xlsx():
    p = os.environ.get("COCO_TEST_XLSX")
    return p if p and os.path.isfile(p) else None


class _NoSeek:
    """seek/tell 없는 쓰기 스트림 — zipfile이 data descriptor(0x08)를 쓰게 만듦."""

    def __init__(self):
        self.buf = io.BytesIO()

    def write(self, b):
        return self.buf.write(b)

    def flush(self):
        pass


def synth_xlsx(path, descriptor=False, ct_last=False):
    """openpyxl로 만든 xlsx. descriptor=True면 0x08 비트가 있는 ZIP으로,
    ct_last=True면 [Content_Types].xml을 마지막 항목으로 다시 포장."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "기본작업-1"
    ws["A1"] = "코코"
    ws["B2"] = 12345
    wb.create_sheet("계산작업")["A1"] = "=SUM(1,2)"
    bio = io.BytesIO()
    wb.save(bio)
    raw = bio.getvalue()
    if not descriptor and not ct_last:
        with open(path, "wb") as f:
            f.write(raw)
        return path
    with zipfile.ZipFile(io.BytesIO(raw)) as zin:
        items = zin.infolist()
        if ct_last:
            items = [i for i in items if i.filename != sj.CONTENT_TYPES_NAME] + \
                [i for i in items if i.filename == sj.CONTENT_TYPES_NAME]
        sink = _NoSeek() if descriptor else io.BytesIO()
        with zipfile.ZipFile(sink, "w", zipfile.ZIP_DEFLATED) as zout:
            for it in items:
                zout.writestr(it.filename, zin.read(it.filename))
        data = sink.buf.getvalue() if descriptor else sink.getvalue()
    with open(path, "wb") as f:
        f.write(data)
    return path


def entry_report(path):
    with zipfile.ZipFile(path) as zf:
        infos = zf.infolist()
        return {
            "first": infos[0].filename,
            "dd": [i.filename for i in infos if i.flag_bits & 0x08],
            "names": [i.filename for i in infos],
            "testzip": zf.testzip(),
            "ct": zf.read(sj.CONTENT_TYPES_NAME).decode("utf-8"),
            "parts": {i.filename: zf.read(i.filename) for i in infos
                      if i.filename != sj.CONTENT_TYPES_NAME},
        }


def assert_clean_conversion(src, dst, label):
    a, b = entry_report(src), entry_report(dst)
    check(f"{label}: [Content_Types].xml이 첫 항목", b["first"] == sj.CONTENT_TYPES_NAME, b["first"])
    check(f"{label}: data descriptor(0x08) 비트 항목 없음", not b["dd"], str(b["dd"][:3]))
    check(f"{label}: zipfile.testzip 통과", b["testzip"] is None, str(b["testzip"]))
    check(f"{label}: 항목 집합 보존(순서는 CT 우선)", set(a["names"]) == set(b["names"])
          and a["names"] and b["names"][1:] == [n for n in a["names"] if n != sj.CONTENT_TYPES_NAME])
    check(f"{label}: 매크로 content-type으로 교체", sj.XLSM_MAIN_CT in b["ct"] and sj.XLSX_MAIN_CT not in b["ct"])
    check(f"{label}: 나머지 파트 바이트 동일", a["parts"] == b["parts"])
    ok, note = sj.verify_xlsm_copy(dst)
    check(f"{label}: verify_xlsm_copy 통과 + openpyxl 로드", ok and "openpyxl 로드 OK" in note, note)
    import openpyxl
    wb = openpyxl.load_workbook(dst)
    check(f"{label}: 시트 이름 보존", wb.sheetnames == openpyxl.load_workbook(src).sheetnames, str(wb.sheetnames))
    wb.close()


print("1. 실제 문제 파일 변환")
real = real_xlsx()
if real:
    dst = os.path.join(TMP, "실제_풀이.xlsm")
    sj.convert_xlsx_to_xlsm(real, dst)
    assert_clean_conversion(real, dst, "실제 xlsx")
    a = entry_report(real)
    print(f"      원본 {os.path.basename(real)}: {len(a['names'])}항목, 0x08 항목 {len(a['dd'])}개")
else:
    print("      (COCO_TEST_XLSX 미지정 — 합성 파일만 검사)")

print("2. 합성 xlsx: 0x08 비트 · CT 순서 어긋남 → 변환 후 깨끗")
dirty = synth_xlsx(os.path.join(TMP, "dirty.xlsx"), descriptor=True, ct_last=True)
d = entry_report(dirty)
check("합성 원본에 0x08 비트 항목이 있고 CT가 첫 항목이 아님", d["dd"] and d["first"] != sj.CONTENT_TYPES_NAME,
      f"dd={len(d['dd'])} first={d['first']}")
ok0, why0 = sj.verify_xlsm_copy(dirty)
check("verify_xlsm_copy는 더러운 원본을 거부", not ok0, why0)
out = sj.convert_xlsx_to_xlsm(dirty, os.path.join(TMP, "clean.xlsm"))
assert_clean_conversion(dirty, out, "합성 xlsx")
plain = synth_xlsx(os.path.join(TMP, "plain.xlsx"))
assert_clean_conversion(plain, sj.convert_xlsx_to_xlsm(plain, os.path.join(TMP, "plain.xlsm")), "일반 xlsx")

print("3. copy_as_macro_enabled")
try:
    os.remove(sj.STARTUP_LOG_PATH)
except OSError:
    pass
src3 = synth_xlsx(os.path.join(TMP, "세트_문제.xlsx"))
dst3 = sj.copy_as_macro_enabled(src3, os.path.join(TMP, "풀이_세트_1"))
check(".xlsx → 풀이 사본 .xlsm", dst3.endswith(".xlsm") and os.path.isfile(dst3), dst3)
check("사본 검증 통과", sj.verify_xlsm_copy(dst3)[0])
log = open(sj.STARTUP_LOG_PATH, encoding="utf-8").read()
check("시작 로그에 사본 경로·크기·변환 성공 기록", "사본 생성" in log and dst3 in log and "변환 성공" in log
      and f"{os.path.getsize(dst3):,} bytes" in log, log.strip()[-120:])
bogus = os.path.join(TMP, "가짜_문제.xlsx")
with zipfile.ZipFile(bogus, "w") as zf:      # [Content_Types].xml 없는 zip
    zf.writestr("xl/workbook.xml", "<workbook/>")
dstb = sj.copy_as_macro_enabled(bogus, os.path.join(TMP, "풀이_가짜"))
check("변환 불가 파일은 .xlsx 그대로 복사(폴백) + 로그", dstb.endswith(".xlsx") and os.path.isfile(dstb)
      and not os.path.exists(os.path.join(TMP, "풀이_가짜.xlsm"))
      and "변환 실패" in open(sj.STARTUP_LOG_PATH, encoding="utf-8").read())
xlsm_src = os.path.join(TMP, "원본_문제.xlsm")
shutil.copy2(dst3, xlsm_src)
dstm = sj.copy_as_macro_enabled(xlsm_src, os.path.join(TMP, "풀이_원본"))
check(".xlsm 원본은 그대로 복사", dstm.endswith(".xlsm") and open(dstm, "rb").read() == open(xlsm_src, "rb").read())
stem = sj._unique_stem(TMP, "풀이_원본")
check("같은 이름 사본이 있으면 _2 어간", stem.endswith("풀이_원본_2"), stem)

print("4. Excel 실행 헬퍼")
check("command 파싱: 따옴표 경로 + /dde",
      sj._exe_from_command('"C:\\Program Files\\Microsoft Office\\root\\Office16\\EXCEL.EXE" /dde')
      == "C:\\Program Files\\Microsoft Office\\root\\Office16\\EXCEL.EXE")
check("command 파싱: 따옴표 없는 경로 + \"%1\"",
      sj._exe_from_command('C:\\Office\\EXCEL.EXE "%1"') == "C:\\Office\\EXCEL.EXE")
check("command 파싱: 빈 값 → None", sj._exe_from_command("") is None and sj._exe_from_command(None) is None)
if sys.platform != "win32":
    check("Windows 아님: find_excel_exe None / excel_running None",
          sj.find_excel_exe() is None and sj.excel_running() is None)
    _env = {k: os.environ.pop(k) for k in ("ProgramFiles", "ProgramFiles(x86)", "ProgramW6432", "LOCALAPPDATA")
            if k in os.environ}
    try:
        check("흔한 경로 후보: 환경변수 없으면 비어 있음", sj._excel_common_paths() == [])
        os.environ["ProgramFiles"] = "/pf"
        cands = sj._excel_common_paths()
        check("흔한 경로 후보: Program Files 아래 Office16/15 EXCEL.EXE 6개",
              len(cands) == 6 and all(c.endswith("EXCEL.EXE") and c.startswith("/pf") for c in cands)
              and any("Office16" in c for c in cands), str(cands[:2]))
    finally:
        os.environ.pop("ProgramFiles", None)
        os.environ.update(_env)
lockdir = os.path.join(TMP, "lock")
os.makedirs(lockdir)
wb_path = os.path.join(lockdir, "풀이_2024 A형_20260906_1030.xlsm")
other = os.path.join(lockdir, "풀이_2024 B형_20260906_1030.xlsm")
for p in (wb_path, other):
    open(p, "wb").close()
check("잠금 파일 없음 → None", sj.workbook_lock_file(wb_path) is None)
full_lock = os.path.join(lockdir, "~$풀이_2024 A형_20260906_1030.xlsm")
open(full_lock, "wb").close()
check("'~$전체이름' 잠금 파일 감지", sj.workbook_lock_file(wb_path) == full_lock)
check("다른 파일(B형)의 잠금으로 오인하지 않음", sj.workbook_lock_file(other) is None)
os.remove(full_lock)
short_lock = os.path.join(lockdir, "~$이_2024 A형_20260906_1030.xlsm")
open(short_lock, "wb").close()
check("앞 글자가 잘린 잠금 파일도 감지", sj.workbook_lock_file(wb_path) == short_lock)
status, detail = sj.check_workbook_open(wb_path)
check("잠금 파일 있으면 'open'", status == "open" and "잠금 파일" in detail, detail)
os.remove(short_lock)


class FakeProc:
    def __init__(self, rc):
        self.rc = rc

    def poll(self):
        return self.rc


if sys.platform != "win32":
    check("Windows 아님 + 잠금 없음 → unknown", sj.check_workbook_open(wb_path)[0] == "unknown")
    check("직접 실행 프로세스가 0이 아닌 코드로 끝남 → closed",
          sj.check_workbook_open(wb_path, FakeProc(1))[0] == "closed")
    check("종료 코드 0(기존 Excel에 파일 전달)만으로는 닫힘으로 보지 않음",
          sj.check_workbook_open(wb_path, FakeProc(0))[0] == "unknown")
    _orig_open = sj.open_file
    sj.open_file = lambda p: (True, "")          # 실제로 프로그램을 띄우지 않음
    try:
        ok, method, err, proc = sj.open_workbook(wb_path)
    finally:
        sj.open_file = _orig_open
    check("open_workbook: Windows 외에서는 기본 프로그램 폴백 (성공, 방법, 오류, proc=None)",
          ok is True and method == "기본 프로그램" and err == "" and proc is None, str((ok, method, err, proc)))
check("상수: EXCEL 확인 지연 5초, PDF 후 1초", sj.EXCEL_CHECK_DELAY_MS == 5000 and sj.EXCEL_OPEN_DELAY_MS == 1000)

print()
print(f"사본 변환·Excel 헬퍼 테스트 {N}건 전부 통과")
