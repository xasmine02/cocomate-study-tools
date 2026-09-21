# -*- coding: utf-8 -*-
"""grade.py 오탐 회귀 스위트 — 실사용 오탐 신고·이의제기의 재현 테스트.

실행: python 회귀테스트.py   (grade.py와 같은 폴더, openpyxl 필요)
모든 테스트는 합성 워크북만 사용하며 외부 파일에 의존하지 않습니다.

다만 합성 = openpyxl 이 쓰는 마크업만 쓴다는 뜻은 아닙니다. openpyxl 은
단추 캡션에 <br/> 를 붙이지 않고, 열 너비에 5픽셀 여백을 더하지 않고,
모듈 스트림을 미니 스트림에 넣지 않습니다 — 그래서 '우리가 만든 세계'
안에서만 참인 검증이 되고 실물에서 터지는 버그를 하나도 못 잡았습니다.
28장의 픽스처는 **실제 Excel 이 저장하는 마크업 형태**를 본떠 만듭니다
(저작권 자료의 내용이 아니라 형태만 본뜹니다). 관찰한 형태는
이의제기_로그.md 의 "Excel 이 실제로 저장하는 형태" 표에 정리돼 있습니다.

수록된 오탐 사례:
  1. 병합 하위 셀 서식 저장 위치 차이 (앵커 기준 비교)
  2. 서식 지시 뒤섞임 (지시 단위 공간 클러스터링·범위 표기·속성 비교)
  3. 인접 셀에 분산 저장된 테두리 (edge 합성 비교)
  4. 테마 색 vs RGB 저장 차이 (테마 해석 후 색 비교, 셀 스타일명 비교 제거)
  5. 지시 범위(diff) 밖 셀의 잔여 서식 (감점 금지)
  6. 자동 조정 행 높이 (기본 제외, key format_checks 명시 시만 채점)
  7. 고급 필터 조건 열/행 배치 차이 (의미 동치 비교)
  8. 적대적 감사 케이스 (회귀케이스/ 36건)
  9. v1.9.0 서식 정밀화: 클러스터 시트 전체 뭉침(행 간격 2행 이상 분리 +
     5개 초과 시 배점 비례 분할), 정답 빈 셀 글꼴/표시형식 오탐 제외
     (빈 셀 채우기는 유지), 속성 행 셀/범위 표기, '확인할 위치' 구분자,
     거의 빈 풀이(일치율 10% 미만) 경고 배너
  10. v2.1.0 리포트 이의제기 **제거**: 카드별 [이의제기] 버튼/입력란, 상단
     고정 바(카운트+복사하기), 폴백 모달, 안내 카드, 관련 JS·CSS·저장소
     키가 리포트에서 사라졌고 오답노트·취약점 진단·해설은 그대로인지
  11. v2.0.1 이의제기 3건(2024 A형) + 문제지 검증 발견 1건(2024 상시 2회):
     ① 기본작업-1 정답 셀 표시 형식 기준 같은 값(0.00% → 57.14%) 정답 처리
     ② 기본작업-3 조건부 서식 셀 단위 실효 규칙 비교(범위·$고정·서식)와
        정답/내 규칙 나란히 표기 ③ 계산작업 셀 상태(빈 셀/수식만 있음)
        표기·학생 수식 전문·문제 수/배점·문제 파일 오염 경고
     ④ 기본작업-2 빈 셀 맞춤 잔여 서식(세로 맞춤) 채점 제외,
        '선택 영역의 가운데로'는 유지
  12. v2.0.2 이의제기 5건(2026 상시 복원 1회, 공식 문제지 PDF 기준):
     ① 세트 기대값 exclude(정답 잔여 서식 F2 맞춤 제외) ② 정의된 이름
        내 답 실제 정의 표기·영역 수 차이 설명 ③ 셀 값 기준 규칙 연산자
        표기('셀 값 >= 0.7')와 '비교 연산자 불일치' 진단 ④ 계산작업 자료 셀
        (정답이 수식 아님) 제외·문제 파일 오염(다른 수식) 경고 ⑤ 차트 축
        기본/보조 역할 비교, 정답 '자동' 속성 제외, 최소 0 = 자동 0 동치,
        세트 기대값 chart.axes
  13. v2.0.3 이의제기 4건(코코 모의고사 1회, 자체 제작 문제지 기준):
     ① Excel [표준 색] 파랑 = #0070C0 · '병합하고 가운데 맞춤'은 가로 가운데만
        (정답 파일에 세로 맞춤 잔여 서식이 있으면 오탐이 남는 것까지 재현)
     ② 표시 형식은 코드가 달라도 그 값의 화면 표시가 같으면 정답
        (정수 범위 0"개" ↔ G/표준"개"), 리포트는 원문 코드 표기
     ③ 비교 연산자 좌우 교환 동치(A>=B ↔ B<=A)와 $ 고정 오답 구분,
        조건부 서식 해설을 이 문제의 범위·기준 열로 생성
     ④ 매크로: 합성 vbaProject.bin(OLE/CFB + MS-OVBA)에서 매크로 이름·범위·
        동작을 읽어 실행 결과가 없어도 인정, 단추 앵커/텍스트/연결 매크로와
        정답/내 답 서식 값 표기, 읽을 수 없으면 '판정 불가: 이유'
     ⑤ (선택) 모의고사/ 세트가 있으면 정답 파일 자기 채점 100점 확인
  14. v2.0.4 이의제기 2건(2024 상시 2회, 코코 복원 문제지 기준):
     ① 계산작업 3번 — 계산값이 없으면 LibreOffice로 재계산해 결과값으로
        비교(어순만 바꾼 식은 만점), 단 결과가 같아도 참조하는 셀이 다르면
        오답 유지 + 셀 단위 근거('정답은 D4, 내 답은 C4'), 재계산을 끄면
        조용히 수식 문자열 비교로 내려가고 그 사실을 리포트에 표기
     ② 매크로 서식 — 다중 영역 Range("B5:F5,B6:B10") 인식, 두 범위 중
        한 범위만 적용했을 때 맞은 범위/틀린 범위 함께 표기, 글꼴 색을
        '파랑(#0070C0)' / '자동(검정)'으로 정밀 표기, 같은 매크로가 있으나
        범위만 다르면 '적용 범위가 지시와 다릅니다', 도형(육각형)의
        텍스트·연결 매크로·도형 종류 대조
  15. v2.0.8 모르면 감점 금지 (UNKNOWN 보류):
     ① 추출 실패를 '없음'과 구분하는 UNKNOWN 센티널 — cfb_streams,
        vba_module_sources/vba_units, form_controls, drawing_shapes,
        chart_features, cf_rules_detail, defined_names_for_sheet,
        fmt_signature/merge_maps/cell_raw_m, sheet_xml/sheet_rel_targets
     ② 채점부는 UNKNOWN 을 만나면 감점하지 않고 '확인 필요'로 보류하고
        그 배점을 만점에서 뺀다 (감점도 가점도 아님)
     ③ 리포트: 총점 옆 '확인 필요 N건', 본문 [확인 필요] 섹션(무엇을 왜),
        점수 계산식 명시, JSON 에 review_* 추가(기존 키 유지)
     ④ 실제 Excel 마크업 픽스처: VML(<br/> 줄바꿈, <x:Anchor> 들여쓰기,
        [0]! 접두 매크로, <v:shapetype> 선행), 시트 <controls>(mc:
        AlternateContent + <xdr:col> 접두사 + controlPr macro), rels 의
        '../' 상대 경로, drawing 의 hidden 단추 그림자, 미니 스트림
        vbaProject.bin, 5픽셀 여백이 붙은 열 너비, 예약 이름(_xlnm./_xleta.)
     ⑤ 미탐 방지: 읽을 수 있는데 틀린 답은 여전히 감점
"""

import importlib.util
import os
import sys
import tempfile

import openpyxl
from openpyxl.styles import Alignment, Border, Color, Font, NamedStyle, \
    PatternFill, Side

BASE = os.path.dirname(os.path.abspath(__file__))
spec = importlib.util.spec_from_file_location("g", os.path.join(BASE, "grade.py"))
g = importlib.util.module_from_spec(spec)
spec.loader.exec_module(g)

TMP = tempfile.mkdtemp()
THIN = Side(style="thin")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def wb_new(sheet):
    wb = openpyxl.Workbook()
    wb.active.title = sheet
    return wb


def save(wb, name):
    p = os.path.join(TMP, name)
    wb.save(p)
    return p


def grade(prob, ans, stu, key=None):
    res, score, notes, _ = g.run_grading(prob, ans, stu, g.normalize_key(key or {}))
    return res[0], score, notes


def check(no, desc, cond, extra=""):
    status = "OK" if cond else "FAIL"
    print(f"  [{no}] {desc}: {status}" + (f" ({extra})" if extra else ""))
    if not cond:
        raise AssertionError(f"오탐 회귀 {no} 실패: {desc} {extra}")


# ---------------------------------------------------------------------------
print("1. 병합 하위 셀 서식 저장 위치 차이")
F16 = Font(name="굴림체", size=16, bold=True, italic=True, underline="double")


def wb_merge_case(merged, font_cells):
    wb = wb_new("기본작업-2")
    ws = wb.active
    ws["B2"] = "제목"
    for r in range(4, 8):
        for c in range(2, 8):
            ws.cell(r, c).value = r * 10 + c
    for co in font_cells:
        ws[co].font = F16
    if merged:
        ws.merge_cells("B2:G2")
    return wb


p1 = save(wb_merge_case(False, []), "1문제.xlsx")
a1 = save(wb_merge_case(True, ["B2", "C2", "D2", "E2", "F2", "G2"]), "1정답.xlsx")
s1 = save(wb_merge_case(True, ["B2"]), "1학생.xlsx")
r, score, _ = grade(p1, a1, s1)
check(1, "앵커에만 글꼴 저장한 학생 만점", r.earned == r.alloc, str(r.details))
s1b = save(wb_merge_case(False, ["B2"]), "1학생b.xlsx")
r, _, _ = grade(p1, a1, s1b)
props1b = str([c.get("props") for c in r.wrong])
check(1, "병합만 안 한 학생: 병합만 지적(글꼴 오탐 없음)",
      r.earned < r.alloc and "병합" in props1b and "글꼴" not in props1b,
      props1b[:120])

# ---------------------------------------------------------------------------
print("2. 서식 지시 뒤섞임 (클러스터 분리 + 범위 표기 + 속성 비교)")


def wb_cluster_case(apply_cd, apply_h):
    wb = wb_new("기본작업-2")
    ws = wb.active
    for r in range(5, 12):
        ws.cell(r, 3).value = f"이름{r}"
        ws.cell(r, 4).value = r
        ws.cell(r, 8).value = 45000 + r
    if apply_cd:
        for r in range(5, 12):
            ws.cell(r, 3).alignment = Alignment(horizontal="distributed")
        ws.merge_cells("D5:D7")
        ws.merge_cells("D8:D11")
        ws["D5"].alignment = Alignment(horizontal="center", vertical="center")
        ws["D8"].alignment = Alignment(horizontal="center", vertical="center")
    if apply_h:
        for r in range(6, 12):
            ws.cell(r, 8).number_format = 'm/d(aaa)'
    return wb


p2 = save(wb_cluster_case(False, False), "2문제.xlsx")
a2 = save(wb_cluster_case(True, True), "2정답.xlsx")
s2 = save(wb_cluster_case(False, True), "2학생.xlsx")  # H 지시만 수행
r, _, _ = grade(p2, a2, s2)
fmt_cards = [c for c in r.wrong if c["kind"] == "format"]
check(2, "실패 카드 1장(C/D 지시)만 생성", len(fmt_cards) == 1,
      str([c["label"] for c in r.wrong]))
card = fmt_cards[0]
chips = " ".join(c["coord"] for c in card["cells"])
check(2, "위치가 범위 표기(C5:C11)이고 H셀 미포함",
      "C5:C11" in chips and "H" not in chips, chips)
props_txt = str(card["props"])
check(2, "속성 비교 '균등 분할' 한국어 표기",
      "가로 맞춤" in props_txt and "균등 분할" in props_txt, props_txt[:120])
check(2, "파생 힌트(균등 분할 경로)", "균등 분할" in (card["hint"] or ""),
      card["hint"])
s2b = save(wb_cluster_case(True, True), "2학생b.xlsx")
r, score, _ = grade(p2, a2, s2b)
check(2, "전부 수행하면 만점", r.earned == r.alloc, str(r.details))

# ---------------------------------------------------------------------------
print("3. 인접 셀 분산 저장 테두리 (edge 합성)")


def wb_border_case(mode):
    wb = wb_new("기본작업-2")
    ws = wb.active
    for r in range(4, 12):
        for c in range(2, 9):
            ws.cell(r, c).value = r * 100 + c
    if mode == "none":
        return wb
    for r in range(4, 12):
        for c in range(2, 9):
            if mode == "scatter" and (r, c) in ((5, 7), (5, 8)):
                continue  # G5/H5 자체 저장 없음 - 이웃에 의존
            ws.cell(r, c).border = BOX
    if mode == "scatter":
        # 이웃 선언으로 G5/H5의 모든 선 커버
        ws.cell(5, 8).border = Border(left=THIN, right=THIN)  # H5: G5|H5 선 + 표 오른쪽
    if mode == "hole":
        ws.cell(11, 8).border = Border(left=THIN, right=THIN, top=THIN)  # H11 아래 없음
    return wb


p3 = save(wb_border_case("none"), "3문제.xlsx")
a3 = save(wb_border_case("full"), "3정답.xlsx")
s3 = save(wb_border_case("scatter"), "3학생.xlsx")
r, _, _ = grade(p3, a3, s3)
check(3, "빈 셀이 이웃 저장에 의존해도 테두리 통과", r.earned == r.alloc,
      str(r.details))
s3b = save(wb_border_case("hole"), "3학생b.xlsx")
r, _, _ = grade(p3, a3, s3b)
border_cards = [c for c in r.wrong if "테두리" in str(c.get("props"))]
check(3, "실제 빠진 선(H11 아래쪽)만 정확히 감점",
      r.earned < r.alloc and border_cards
      and "H11 아래쪽 선 없음" in str(border_cards[0]["props"]),
      str(border_cards[0]["props"])[:150] if border_cards else str(r.details))

# ---------------------------------------------------------------------------
print("4. 테마 색 vs RGB 저장 차이 + 셀 스타일명 비교 제거")


def wb_color_case(mode, accent4_rgb="FFC000"):
    wb = wb_new("기본작업-2")
    ws = wb.active
    for c in range(2, 9):
        ws.cell(4, c).value = f"머리글{c}"
    if mode == "none":
        return wb
    for c in range(2, 9):
        cell = ws.cell(4, c)
        if mode == "theme":
            cell.fill = PatternFill("solid", fgColor=Color(theme=7))
            cell.font = Font(color=Color(theme=0), bold=True)
        elif mode == "rgb":
            cell.fill = PatternFill("solid", fgColor="FF" + accent4_rgb)
            cell.font = Font(color="FFFFFFFF", bold=True)
        elif mode == "named":
            if "머리글st" not in wb.named_styles:
                ns = NamedStyle(name="머리글st")
                ns.fill = PatternFill("solid", fgColor="FF" + accent4_rgb)
                ns.font = Font(color="FFFFFFFF", bold=True)
                wb.add_named_style(ns)
            cell.style = "머리글st"
        elif mode == "wrongcolor":
            cell.fill = PatternFill("solid", fgColor="FF00B050")  # 초록
            cell.font = Font(color="FFFFFFFF", bold=True)
    return wb


p4 = save(wb_color_case("none"), "4문제.xlsx")
a4 = save(wb_color_case("theme"), "4정답.xlsx")
# 정답 파일이 실제로 쓰는 테마 팔레트에서 accent4/배경1 RGB를 얻어
# "같은 색을 RGB로 저장한 학생" 파일을 만든다 (저장 방식만 다른 동일 색)
_a4wb = openpyxl.load_workbook(a4)
_pal = g._theme_palette(_a4wb)
ACCENT4 = _pal[7]
check(4, "배경1(theme 0) = 흰색 환산", _pal[0] == "FFFFFF", _pal[0])
for mode, desc in (("rgb", "같은 색 RGB 저장"), ("named", "named style 지정")):
    s4 = save(wb_color_case(mode, ACCENT4), f"4학생_{mode}.xlsx")
    r, _, _ = grade(p4, a4, s4)
    check(4, f"{desc} 학생 통과", r.earned == r.alloc, str(r.details))
s4w = save(wb_color_case("wrongcolor"), "4학생_wrong.xlsx")
r, _, _ = grade(p4, a4, s4w)
props = str([c.get("props") for c in r.wrong])
check(4, "실제 다른 색(초록)은 감점 + 색 이름 표기",
      r.earned < r.alloc and "초록" in props, props[:150])
# 테마 tint 환산 수동 대조 2케이스
pal_probe = ["FFC000"]
check(4, "tint 환산: accent4 25% 어둡게 ≈ #BF8F00",
      g._quant_hex(g._apply_tint("FFC000", -0.25)) == g._quant_hex("BF8F00"),
      g._apply_tint("FFC000", -0.25))
check(4, "tint 환산: 흰색(배경1) tint 0 = FFFFFF",
      g._apply_tint("FFFFFF", 0.0) == "FFFFFF")

# ---------------------------------------------------------------------------
print("5. 지시 범위(diff) 밖 잔여 서식 감점 금지")


def wb_outside_case(role):
    wb = wb_new("기본작업-2")
    ws = wb.active
    ws["B2"] = "제목"
    for r in range(6, 12):
        ws.cell(r, 8).value = 45000 + r
    if role in ("ans", "stu"):
        ws["B2"].font = Font(bold=True, size=14)   # 유일한 지시
    if role == "stu":
        ws["H5"].number_format = "mm-dd-yy"        # 지시 밖 잔여 서식
    return wb


p5 = save(wb_outside_case("prob"), "5문제.xlsx")
a5 = save(wb_outside_case("ans"), "5정답.xlsx")
s5 = save(wb_outside_case("stu"), "5학생.xlsx")
r, _, _ = grade(p5, a5, s5)
check(5, "diff 밖 H5 잔여 서식은 감점/카드 없음",
      r.earned == r.alloc and not any("H5" in str(c) for c in r.wrong),
      str(r.details))

# ---------------------------------------------------------------------------
print("6. 자동 조정 행 높이 (기본 제외, key 명시 시만 채점)")


def wb_rowh_case(h4):
    wb = wb_new("기본작업-2")
    ws = wb.active
    ws["B2"] = "제목"
    ws["B2"].font = Font(bold=True)
    ws["B4"] = "머리글"
    if h4 is not None:
        ws.row_dimensions[4].height = h4
    return wb


p6 = save(wb_rowh_case(None), "6문제.xlsx")
a6 = save(wb_rowh_case(33.75), "6정답.xlsx")
s6 = save(wb_rowh_case(30.0), "6학생.xlsx")
a6wb = openpyxl.load_workbook(a6)
a6wb.active["B2"].font = Font(bold=True)
a6wb.save(a6)
s6wb = openpyxl.load_workbook(s6)
s6wb.active["B2"].font = Font(bold=True)
s6wb.save(s6)
r, _, notes = grade(p6, a6, s6)
check(6, "행 높이 차이 기본 감점 0 + 참고 노트",
      r.earned == r.alloc and any("행 높이" in n for n in r.notes),
      str(r.notes))
key6 = {"format_checks": {"기본작업-2": [
    {"no": 9, "range": "4:4", "check": {"row_height": 30}}]}}
r, _, _ = grade(p6, a6, s6, key=key6)
check(6, "key 명시 행높이 30 = 학생 30 통과", r.earned == r.alloc,
      str(r.details))
s6b = save(wb_rowh_case(20.0), "6학생b.xlsx")
s6bwb = openpyxl.load_workbook(s6b)
s6bwb.active["B2"].font = Font(bold=True)
s6bwb.save(s6b)
r, _, _ = grade(p6, a6, s6b, key=key6)
check(6, "key 명시 행높이 미달(20) 감점 + '4행' 구분 표기",
      r.earned < r.alloc and any("4행" in d for d in r.details),
      str(r.details))
key6b = {"format_checks": {"기본작업-2": [
    {"no": 9, "range": "1행", "check": {"row_height": 30}},
    {"no": 9, "range": "A열", "check": {"col_width": 12}}]}}
a6c = openpyxl.load_workbook(a6)
a6c.active.row_dimensions[1].height = 30.0
a6c.active.column_dimensions["A"].width = 12.0
a6cp = save(a6c, "6정답c.xlsx")
s6c = openpyxl.load_workbook(s6)
s6c.active.row_dimensions[1].height = 30.2
s6c.active.column_dimensions["A"].width = 8.0
s6cp = save(s6c, "6학생c.xlsx")
r, _, _ = grade(p6, a6cp, s6cp, key=key6b)
check(6, "key 열너비 미달만 감점('A열' 표기), 행높이 ±0.5 허용",
      any("A열" in d for d in r.details)
      and not any("[행 높이]" in d for d in r.details), str(r.details))

# ---------------------------------------------------------------------------
print("7. 고급 필터 조건/결과 의미 동치")

DATA = [("이름", "성별", "나이"), ("kate", "여성", 24), ("brad", "남성", 31),
        ("ann", "여성", 29), ("mark", "남성", 27), ("sarah", "여성", 35)]


def wb_filter_case(cond_layout, result_rows=None, cond2=None):
    wb = wb_new("기본작업-3")
    ws = wb.active
    for i, row in enumerate(DATA, start=1):
        for j, v in enumerate(row, start=1):
            ws.cell(i, j).value = v
    if cond_layout:
        for (r, c, v) in cond_layout:
            ws.cell(r, c).value = v
    if cond2:
        for (r, c, v) in cond2:
            ws.cell(r, c).value = v
    if result_rows:
        for i, row in enumerate(result_rows):
            for j, v in enumerate(row):
                ws.cell(20 + i, 2 + j).value = v
    return wb


RESULT = [("이름", "나이"), ("kate", 24), ("ann", 29), ("sarah", 35)]
COND_A = [(16, 2, "성별"), (16, 3, "이름"), (17, 2, "<>남성"), (17, 3, "*a*")]
COND_SWAP = [(16, 2, "이름"), (16, 3, "성별"), (17, 2, "*a*"), (17, 3, "<>남성")]
COND_WRONG = [(16, 2, "성별"), (16, 3, "이름"), (17, 2, "남성"), (17, 3, "*a*")]

p7 = save(wb_filter_case(None), "7문제.xlsx")
a7 = save(wb_filter_case(COND_A, RESULT), "7정답.xlsx")
s7 = save(wb_filter_case(COND_SWAP, RESULT), "7학생_swap.xlsx")
r, _, _ = grade(p7, a7, s7)
check(7, "조건 범위 열 스왑 -> 만점", r.earned == r.alloc, str(r.details))
s7w = save(wb_filter_case(COND_WRONG, RESULT), "7학생_wrong.xlsx")
r, _, _ = grade(p7, a7, s7w)
det = " ".join(r.details)
check(7, "진짜 오답은 (필드,조건) 짝 메시지로 감점",
      r.earned < r.alloc and "성별" in det and "<>남성" in det, det[:150])
# OR 두 행 조건: 행 순서 스왑
COND_OR = [(16, 2, "성별"), (16, 3, "나이"), (17, 2, "<>남성"), (18, 3, ">=30")]
COND_OR_SWAP = [(16, 2, "성별"), (16, 3, "나이"), (17, 3, ">=30"), (18, 2, "<>남성")]
a7o = save(wb_filter_case(COND_OR, RESULT), "7정답_or.xlsx")
s7o = save(wb_filter_case(COND_OR_SWAP, RESULT), "7학생_or.xlsx")
r, _, _ = grade(p7, a7o, s7o)
check(7, "OR 조건 행 순서 스왑 -> 통과", r.earned == r.alloc, str(r.details))
# 결과 표 행 순서 스왑 (집합 동치 + 무감점 노트)
RESULT_SWAP = [RESULT[0], RESULT[2], RESULT[1], RESULT[3]]
s7r = save(wb_filter_case(COND_A, RESULT_SWAP), "7학생_rows.xlsx")
r, _, _ = grade(p7, a7, s7r)
check(7, "결과 행 순서 스왑 -> 통과 + 무감점 노트",
      r.earned == r.alloc and any("행 순서" in n for n in r.notes),
      str(r.notes))

print()
print("오탐 회귀 스위트 7건 전부 통과")

# ---------------------------------------------------------------------------
# 8. 적대적 감사 케이스 (회귀케이스/ 36건) — v1.5.0에서 병합
# ---------------------------------------------------------------------------
print("8. 적대적 감사 케이스 (회귀케이스/)")
import subprocess

CASES = os.path.join(BASE, "회귀케이스")
if not os.path.isdir(CASES):
    print("  회귀케이스/ 폴더 없음 — 건너뜀")
else:
    import json
    mismatch = []
    total = 0
    for cid in sorted(os.listdir(CASES)):
        d = os.path.join(CASES, cid)
        meta_p = os.path.join(d, "expected.json")
        if not os.path.isdir(d) or not os.path.isfile(meta_p):
            continue
        with open(meta_p, encoding="utf-8") as f:
            meta = json.load(f)
        total += 1
        try:
            if meta.get("cli"):
                sub = next(os.path.join(d, n) for n in os.listdir(d)
                           if os.path.isdir(os.path.join(d, n)))
                files = sorted(os.listdir(sub))
                prob = next(f2 for f2 in files if "문제" in f2)
                ans = next(f2 for f2 in files if "정답" in f2)
                stu = next(f2 for f2 in files if "풀이" in f2)
                oj = os.path.join(sub, "결과.json")
                cp = subprocess.run(
                    [sys.executable, os.path.join(BASE, "grade.py"),
                     "--problem", os.path.join(sub, prob),
                     "--answer", os.path.join(sub, ans),
                     "--student", os.path.join(sub, stu),
                     "--json", oj, "--html", os.path.join(sub, "결과.html")],
                    capture_output=True, text=True)
                if cp.returncode != 0:
                    raise RuntimeError(f"exit={cp.returncode}")
                with open(oj, encoding="utf-8") as f:
                    data = json.load(f)
                lost = [s2 for s2 in data["sheets"]
                        if s2["earned"] < s2["alloc"]]
            else:
                res, _sc, _n, _ = g.run_grading(
                    os.path.join(d, "문제.xlsx"), os.path.join(d, "정답.xlsx"),
                    os.path.join(d, "학생.xlsx"),
                    g.normalize_key(meta.get("key") or {}))   # 세트 기대값 케이스
                lost = [r2 for r2 in res if r2.earned < r2.alloc]
            actual = "deduction" if lost else "no_deduction"
            if actual != meta["correct_behavior"]:
                mismatch.append(f"{cid}: {actual} != "
                                f"{meta['correct_behavior']}")
        except Exception as e:
            mismatch.append(f"{cid}: ERROR {e}")
    if mismatch:
        for mm in mismatch:
            print("  [FAIL]", mm)
        raise AssertionError(f"감사 케이스 {len(mismatch)}건 불일치")
    print(f"  감사 케이스 {total}건 전부 correct_behavior 일치")

# ---------------------------------------------------------------------------
# 9. v1.9.0 서식 채점 정밀화 (실사용 오탐 2건 + 리포트 경고 배너)
# ---------------------------------------------------------------------------
print("9. v1.9.0 서식 정밀화 (클러스터 행 간격·빈 셀 글꼴·경고 배너)")

F_TITLE = Font(name="굴림체", size=16, bold=True)
CC = Alignment(horizontal="centerContinuous")
GOLD = PatternFill("solid", fgColor="FFFFC000")


def wb_v19(title_font_cols=(), cc=False, numfmt=False, fill=False):
    """2행 제목 + 4~11행 표. 지시: 제목 글꼴/선택 영역의 가운데로,
    D열 표시형식, F열(값이 빈 셀) 채우기."""
    wb = wb_new("기본작업-2")
    ws = wb.active
    ws["B2"] = "코코 상사 급여 현황"
    for r in range(4, 12):
        ws.cell(r, 2).value = f"이름{r}"
        ws.cell(r, 4).value = 45000 + r
    for c in title_font_cols:
        ws.cell(2, c).font = F_TITLE
    if cc:
        for c in range(2, 9):
            ws.cell(2, c).alignment = CC
    if numfmt:
        for r in range(4, 12):
            ws.cell(r, 4).number_format = "#,##0"
    if fill:
        for r in range(4, 12):
            ws.cell(r, 6).fill = GOLD  # F4:F11 - 값 없는 빈 셀 채우기
    return wb


ALLC = tuple(range(2, 9))
p9 = save(wb_v19(), "9문제.xlsx")
# 정답 제작자는 B2 제목 지시 글꼴을 B2:H2 전체에 걸었음 (빈 셀 포함)
a9 = save(wb_v19(ALLC, cc=True, numfmt=True, fill=True), "9정답.xlsx")

# (a) 지시대로 B2에만 글꼴 + 전체 centerContinuous 적용한 학생 -> 감점 0
s9a = save(wb_v19((2,), cc=True, numfmt=True, fill=True), "9학생a.xlsx")
r, _, _ = grade(p9, a9, s9a)
check(9, "정답 B2:H2 글꼴, 학생 B2만 글꼴(빈 셀 글꼴 오탐 없음) 만점",
      r.earned == r.alloc and not r.wrong, str(r.details))

# (b) 제목 지시만 빠뜨린 학생 -> 2행 클러스터만 감점 (표와 절대 병합 금지)
s9b = save(wb_v19((), cc=False, numfmt=True, fill=True), "9학생b.xlsx")
r, _, _ = grade(p9, a9, s9b)
fmt9 = [c for c in r.wrong if c["kind"] == "format"]
check(9, "2행 제목 지시만 실패: 카드 1장, 4~11행 표 셀 미포함",
      len(fmt9) == 1 and not any(ch["coord"] not in ("B2", "B2:H2")
                                 for ch in fmt9[0]["cells"]),
      str([(c["label"], [ch["coord"] for ch in c["cells"]]) for c in fmt9]))
check(9, "감점은 표 클러스터 몫 제외한 5점(배점 절반)",
      abs(fmt9[0]["lost"] - 5.0) < 0.05 and r.earned == r.alloc - 5,
      f"lost={fmt9[0]['lost']} earned={r.earned}")
names9 = [p["name"] for p in fmt9[0]["props"]]
check(9, "속성 행에 해당 셀/범위 표기: 글꼴(B2)·가로 맞춤(B2:H2)",
      "글꼴(B2)" in names9 and "가로 맞춤(B2:H2)" in names9, str(names9))

# 두 지시 모두 실패 -> 클러스터 2개로 분리(행 간격 2행), 뒤섞임 없음
r, _, _ = grade(p9, a9, p9)
fmt9c = [c for c in r.wrong if c["kind"] == "format"]
check(9, "2행 제목·4~11행 표는 별도 클러스터(카드 2장)", len(fmt9c) == 2,
      str([c["label"] for c in fmt9c]))
title_card = next(c for c in fmt9c
                  if any(ch["coord"].startswith("B2") for ch in c["cells"]))
table_card = next(c for c in fmt9c if c is not title_card)
check(9, "제목 카드에 4행 이후 셀 인용 없음(뒤섞임 방지)",
      not any(p["name"] for p in title_card["props"]
              if any(f"{col}{row}" in p["name"] for col in "BCDEFGH"
                     for row in range(4, 12))),
      str([p["name"] for p in title_card["props"]]))
tprops = str(table_card["props"])
check(9, "(c) 값이 빈 F4:F11 채우기 실누락은 여전히 감점",
      "채우기(F4:F11)" in tprops and "황금색" in tprops, tprops[:150])
html_card = g._wrong_card_html(table_card)
check(9, "'확인할 위치' 나열에 ', ' 구분자",
      "</span>, <span" in html_card, html_card[:200])
# 빈 셀 채우기'만' 빠뜨린 학생도 감점되는지 단독 확인
s9c = save(wb_v19(ALLC, cc=True, numfmt=True, fill=False), "9학생c.xlsx")
r, _, _ = grade(p9, a9, s9c)
check(9, "(c) 채우기만 빠뜨린 학생: 빈 셀이어도 감점",
      r.earned < r.alloc and "채우기(F4:F11)" in str(r.wrong),
      f"earned={r.earned}/{r.alloc}")

# 5개 초과 클러스터: 무리하게 합치지 않고 배점 비례 분할 (합계 유지)
def wb_many(bold_rows):
    wb = wb_new("기본작업-2")
    ws = wb.active
    for i, row in enumerate(range(2, 21, 3)):  # 2,5,8,...,20행 (간격 3행)
        ws.cell(row, 2).value = f"항목{i}"
        if row in bold_rows:
            ws.cell(row, 2).font = Font(bold=True)
    return wb


ROWS7 = list(range(2, 21, 3))  # 7개 지시
p9m = save(wb_many(()), "9문제m.xlsx")
a9m = save(wb_many(ROWS7), "9정답m.xlsx")
s9m = save(wb_many(ROWS7[:4]), "9학생m.xlsx")  # 4개만 수행
r, _, _ = grade(p9m, a9m, s9m)
check(9, "7개 지시 강제 병합 없음: 실패 카드 3장 + 비례 분할(10*4/7=6점)",
      len([c for c in r.wrong if c["kind"] == "format"]) == 3
      and r.earned == 6, f"earned={r.earned} cards={len(r.wrong)}")
r, _, _ = grade(p9m, a9m, save(wb_many(ROWS7), "9학생m2.xlsx"))
check(9, "7개 지시 전부 수행 시 합계 10점 유지", r.earned == r.alloc,
      f"earned={r.earned}")

# (d) 거의 빈 풀이(채점 대상 diff 셀 일치율 10% 미만) -> 헤더 경고 배너
BANNER = "풀이가 거의 비어 있습니다"
res9, sc9, notes9, _ = g.run_grading(p7, a7, p7, g.normalize_key({}))
out9 = os.path.join(TMP, "9배너.html")
g.write_html(out9, res9, sc9, notes9, (p7, a7, p7))
with open(out9, encoding="utf-8") as f:
    html9 = f.read()
check(9, "거의 빈 풀이(문제 파일 그대로 채점) -> 경고 배너", BANNER in html9)
check(9, "리포트 헤더 '풀이' 파일명 시각 강조", 'class="file-solve"' in html9)
res9n, sc9n, notes9n, _ = g.run_grading(p7, a7, a7, g.normalize_key({}))
out9n = os.path.join(TMP, "9배너없음.html")
g.write_html(out9n, res9n, sc9n, notes9n, (p7, a7, a7))
with open(out9n, encoding="utf-8") as f:
    html9n = f.read()
check(9, "정상 풀이 -> 경고 배너 없음", BANNER not in html9n)

# ---------------------------------------------------------------------------
# 10. v2.1.0 리포트에서 이의제기 기능 제거 (오픈 소스 전환)
#     — 없앤 것은 "이의를 제기해 전달하는 경로"뿐이다. 오답노트·취약점 진단·
#       해설은 그대로 남아야 한다.
# ---------------------------------------------------------------------------
print("10. v2.1.0 리포트 이의제기 기능 제거 (오답노트·진단·해설은 유지)")
import json as json_mod
import re as re_mod
import shutil

# 오답 카드 2장이 나오는 혼합 리포트 (섹션 9의 p9/a9 재사용, 학생=문제)
res10, sc10, notes10, _ = g.run_grading(p9, a9, p9, g.normalize_key({}))
out10 = os.path.join(TMP, "10리포트.html")
g.write_html(out10, res10, sc10, notes10, (p9, a9, p9))
with open(out10, encoding="utf-8") as f:
    html10 = f.read()
n_cards10 = sum(len(r2.wrong) for r2 in res10)
check(10, "오답 카드는 그대로 그려진다(제거된 건 이의제기 경로뿐)",
      n_cards10 >= 2 and html10.count('class="wc"') == n_cards10
      and "오답노트" in html10,
      f"cards={n_cards10} wc={html10.count(chr(39) + 'class=' + chr(34) + 'wc' + chr(34) + chr(39))}")
check(10, "카드별 [이의제기] 버튼·입력란·[저장]이 없다",
      "appeal-btn" not in html10 and "appeal-text" not in html10
      and "appeal-save" not in html10
      and "무엇이 잘못됐는지 적어주세요" not in html10)
check(10, "상단 고정 바(이의제기 건수 + [이의제기 복사하기])가 없다",
      "appeal-bar" not in html10 and "appealCount" not in html10
      and "이의제기 복사하기" not in html10 and "has-appeal" not in html10)
check(10, "복사 폴백 모달·execCommand 클립보드 코드가 없다",
      "appealModal" not in html10 and "execCommand" not in html10
      and "직접 드래그해 복사" not in html10)
check(10, "이의제기 JS(buildAppealText·APPEAL_META/DATA)와 저장소 키가 없다",
      "buildAppealText" not in html10 and "APPEAL_META" not in html10
      and "APPEAL_DATA" not in html10 and "kocoAppeal" not in html10
      and 'id="appeal-assemble"' not in html10)
check(10, "'이의제기 안내' 카드와 Claude 채팅 붙여넣기 안내가 없다",
      "이의제기 안내" not in html10 and "이의제기" not in html10
      and "Claude" not in html10
      and "\\uc774\\uc758\\uc81c\\uae30" not in html10)
check(10, "이의제기용 CSS 클래스가 리포트에 남아 있지 않다",
      ".appeal" not in html10 and "appeal-guide" not in html10)
check(10, "모듈에서 이의제기 헬퍼가 사라졌다",
      not hasattr(g, "_appeal_payload") and not hasattr(g, "_APPEAL_UI_JS")
      and not hasattr(g, "_APPEAL_ASSEMBLE_JS"))
check(10, "남긴 기능: 오답노트 · 취약점 진단 · 영역별 점수",
      "오답노트" in html10 and "취약점 진단" in html10
      and "영역별 점수" in html10)
# 카드 내용(내 답 ↔ 정답 비교)은 그대로 — 이의제기 텍스트가 담던 정보가
# 리포트 화면에는 남아 있어야 사용자가 직접 판단하고 고칠 수 있다.
card10 = g._wrong_card_html({
    "sheet": "계산작업", "label": "계산 문제 1", "lost": 8,
    "cells": [{"coord": "J3", "got": "#VALUE!", "expected": "본선",
               "got_formula": "=OR(IF())"}],
    "diff_notes": ["중첩 순서가 반대입니다"], "explain": ["① SMALL(...)"],
    "hint": "확인"})
check(10, "카드에 내 답/정답 비교와 해설이 그대로 있다(제거된 건 이의제기뿐)",
      '내 답' in card10 and '정답' in card10 and "#VALUE!" in card10
      and "무엇이 다른가" in card10 and "정확한 풀이" in card10
      and "appeal" not in card10 and "이의제기" not in card10, card10[:160])
check(10, "리포트 본문에도 오답 카드가 실린다(위치 칩 또는 비교 박스)",
      html10.count('class="wc-top"') == n_cards10
      and ("wc-chips" in html10 or 'class="box mine"' in html10))
# 만점 리포트도 정상
res10p, sc10p, notes10p, _ = g.run_grading(p9, a9, a9, g.normalize_key({}))
out10p = os.path.join(TMP, "10만점.html")
g.write_html(out10p, res10p, sc10p, notes10p, (p9, a9, a9))
with open(out10p, encoding="utf-8") as f:
    html10p = f.read()
check(10, "만점 리포트도 이의제기 흔적 없이 정상 생성",
      "틀린 항목이 없습니다" in html10p and "appeal" not in html10p
      and "이의제기" not in html10p)
check(10, "<body> 에 has-appeal 클래스가 붙지 않는다(상단 여백 보정 제거)",
      "<body>" in html10 and "<body>" in html10p, html10[:0])

# ---------------------------------------------------------------------------
# 11. v2.0.1 이의제기 3건(2024 A형) + 문제지 검증 발견 1건(2024 상시 2회)
# ---------------------------------------------------------------------------
print("11. v2.0.1 이의제기 반영 (표시 형식 값·조건부 서식 실효 규칙·셀 상태 표기·"
      "빈 셀 맞춤)")
import re as re11
import zipfile as zip11
from openpyxl.formatting.rule import FormulaRule


def inject_cache11(path, coord, value):
    """수식 셀에 캐시값 <v>를 주입 (Excel 저장 파일 흉내, 단일 시트 전제)."""
    part = "xl/worksheets/sheet1.xml"
    zin = zip11.ZipFile(path)
    items = [(n, zin.read(n)) for n in zin.namelist()]
    zin.close()
    pat = re11.compile(r'(<c r="%s"[^>]*>)(<f>.*?</f>)(<v\s*/>|<v></v>)?' % coord)
    with zip11.ZipFile(path, "w", zip11.ZIP_DEFLATED) as zout:
        for n, data in items:
            if n == part:
                x = data.decode("utf-8")
                x = pat.sub(lambda m: m.group(1) + m.group(2)
                            + f"<v>{value}</v>", x, count=1)
                data = x.encode("utf-8")
            zout.writestr(n, data)


# --- ① 기본작업-1: 정답 셀 표시 형식 기준 같은 값 -> 정답 ---------------------
def wb_b1(values=None, fmt="0.00%"):
    wb = wb_new("기본작업-1")
    ws = wb.active
    for r in range(4, 7):
        ws.cell(r, 6).number_format = fmt      # 문제 파일에 미리 적용된 형식
    for k, v in (values or {}).items():
        ws[k] = v
    return wb


ANS_B1 = {"A3": "학과코드", "B3": "학과명", "C3": "전체 학생수",
          "E3": "정원/전임(겸임)", "A4": "KA-45267", "C4": 140,
          "F4": 0.5, "F5": 4 / 7, "F6": 2 / 3}
p11 = save(wb_b1(), "11문제.xlsx")
a11 = save(wb_b1(ANS_B1), "11정답.xlsx")
s11a = save(wb_b1(dict(ANS_B1, F5=0.5714, F6=0.6667)), "11학생a.xlsx")
r, _, _ = grade(p11, a11, s11a)
check(11, "0.00% 형식 셀에 57.14%/66.67% 입력(0.5714/0.6667) -> 만점",
      r.earned == r.alloc and any("표시 형식 기준" in n for n in r.notes),
      str(r.details) + str(r.notes))
s11b = save(wb_b1(dict(ANS_B1, F5=0.5814)), "11학생b.xlsx")
r, _, _ = grade(p11, a11, s11b)
check(11, "58.14% 입력(표시도 다름) -> 0점 유지", r.earned == 0, str(r.details))
a11g = save(wb_b1(ANS_B1, fmt="General"), "11정답g.xlsx")
p11g = save(wb_b1(fmt="General"), "11문제g.xlsx")
s11g = save(wb_b1(dict(ANS_B1, F5=0.5714, F6=0.6667), fmt="General"),
            "11학생g.xlsx")
r, _, _ = grade(p11g, a11g, s11g)
check(11, "표시 형식 없는(General) 셀은 0.5714 vs 4/7 감점 유지",
      r.earned == 0, str(r.details))
s11c = save(wb_b1(dict(ANS_B1, C3=None, F5=0.5714, F6=0.6667)), "11학생c.xlsx")
r, _, _ = grade(p11, a11, s11c)
cells11 = r.wrong[0]["cells"] if r.wrong else []
check(11, "빈 셀은 '(빈 셀 — 값·수식 없음)'으로 상태 표기",
      r.earned == 0 and cells11 and cells11[0]["coord"] == "C3"
      and cells11[0]["got"] == g.STATE_BLANK
      and cells11[0].get("got_state") == "blank", str(cells11))
card11 = g._wrong_card_html(r.wrong[0])
check(11, "리포트 카드: 셀 위치 · 정답 · 내 답(빈 셀 상태)이 그대로 보인다",
      ">C3<" in card11 and "전체 학생수" in card11
      and g.STATE_BLANK in card11, card11[:200])

# --- ② 기본작업-3: 조건부 서식 셀 단위 실효 규칙 비교 -----------------------
RED = Font(color="FFFF0000")


def wb_cf(rules=()):
    """A3:H3 머리글 + 학번(A4:A18) 표. rules: [(범위, 수식, 글꼴)]"""
    wb = wb_new("기본작업-3")
    ws = wb.active
    for c, h in enumerate(("학번", "이름", "중간", "중간(40)", "기말", "기말(40)",
                           "출석", "합계"), start=1):
        ws.cell(3, c).value = h
    for i, r in enumerate(range(4, 19)):
        ws.cell(r, 1).value = 201900000 + i if i % 3 == 0 else 201700000 + i
        for c in range(2, 9):
            ws.cell(r, c).value = f"v{r}{c}"
    for rng, f, font in rules:
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f], font=font))
    return wb


p12 = save(wb_cf(), "12문제.xlsx")
a12 = save(wb_cf([("A4:H18", 'LEFT($A4,4)="2019"', RED)]), "12정답.xlsx")
r, _, _ = grade(p12, a12, a12)
check(12, "정답과 같은 규칙(A4:H18, $A4) -> 만점", r.earned == r.alloc,
      str(r.details))
s12b = save(wb_cf([("A4:A18", 'LEFT(A4,4)="2019"', RED)]), "12학생b.xlsx")
r, _, _ = grade(p12, a12, s12b)
card12 = r.wrong[0] if r.wrong else {}
dn12 = " ".join(card12.get("diff_notes") or [])
check(12, "A열만 적용($A$4:$A$18) + 상대참조 -> 0점 (행 전체 아님)",
      r.earned == 0 and "적용 범위 불일치" in dn12 and "B4:H18" in dn12,
      dn12[:200])
check(12, "카드에 정답/내 규칙 나란히(적용 범위·수식·적용 서식 props)",
      [p["name"] for p in card12.get("props", [])][:3]
      == ["적용 범위", "수식", "적용 서식"]
      and card12["props"][0]["got"] == "A4:A18"
      and card12["props"][1]["got"] == '=LEFT(A4,4)="2019"'
      and "빨강" in card12["props"][2]["got"], str(card12.get("props")))
check(12, "해설 포인트에 '행 전체 = 표 전체 범위 + $열 고정' 원리",
      "행 전체" in (card12.get("point") or "") and "$A4" in card12.get("point"),
      card12.get("point", "")[:80])
s12c = save(wb_cf([("A4:H18", 'LEFT(A4,4)="2019"', RED)]), "12학생c.xlsx")
r, _, _ = grade(p12, a12, s12c)
dn12c = " ".join(r.wrong[0].get("diff_notes") or []) if r.wrong else ""
check(12, "범위는 맞고 상대참조 A4 -> 0점, '참조 고정($) 방식' 진단",
      r.earned == 0 and "참조 고정" in dn12c, dn12c[:160])
s12g = save(wb_cf([("A4:H18", 'LEFT($A$4,4)="2019"', RED)]), "12학생g.xlsx")
r, _, _ = grade(p12, a12, s12g)
check(12, "절대참조 $A$4 -> 0점, '참조 고정' 진단",
      r.earned == 0 and "참조 고정" in " ".join(r.wrong[0].get("diff_notes")),
      str(r.wrong[0].get("diff_notes"))[:160])
s12d = save(wb_cf([("A4:H10", 'LEFT($A4,4)="2019"', RED),
                   ("A11:H18", 'LEFT($A11,4)="2019"', RED)]), "12학생d.xlsx")
r, _, _ = grade(p12, a12, s12d)
check(12, "두 구간으로 나눠 적용(각 구간 첫 행 기준 $A4/$A11) -> 만점(화면 동일)",
      r.earned == r.alloc, str(r.details))
s12e = save(wb_cf([("A4:H18", 'LEFT($A4,4)="2019"', Font(color="FF0070C0"))]),
            "12학생e.xlsx")
r, _, _ = grade(p12, a12, s12e)
dn12e = " ".join(r.wrong[0].get("diff_notes") or []) if r.wrong else ""
check(12, "규칙은 같고 글꼴 색만 파랑 -> 0점, '적용 서식 불일치' 진단",
      r.earned == 0 and "적용 서식 불일치" in dn12e and "빨강" in dn12e,
      dn12e[:160])
s12f = save(wb_cf(), "12학생f.xlsx")
r, _, _ = grade(p12, a12, s12f)
check(12, "규칙 없음 -> 0점, 내 답 '(규칙 없음)' 표기",
      r.earned == 0 and r.wrong[0]["props"][0]["got"] == "(규칙 없음)"
      and "규칙이 없습니다" in " ".join(r.wrong[0]["diff_notes"]),
      str(r.wrong[0]["props"])[:160])
props12 = {p.get("name"): p for p in r.wrong[0]["props"]}
card12 = g._wrong_card_html(r.wrong[0])
check(12, "리포트 카드에 규칙 3속성(범위·수식·서식)이 모두 표기된다",
      {"적용 범위", "수식", "적용 서식"} <= set(props12)
      and "적용 범위" in card12 and "수식" in card12 and "적용 서식" in card12,
      str(sorted(props12)))

# --- ③ 계산작업: 셀 상태 표기, 문제 수·배점, 문제 파일 오염 ------------------
CALC_COLS = {"D": "A{r}*2", "F": "A{r}+B{r}", "H": "MAX(A{r},B{r})",
             "J": "A{r}-B{r}", "L": "A{r}*B{r}"}


def wb_calc11(cols=CALC_COLS, blank=(), override=None):
    wb = wb_new("계산작업")
    ws = wb.active
    ws["A2"], ws["B2"] = "값1", "값2"
    for r in range(3, 10):
        ws[f"A{r}"], ws[f"B{r}"] = r * 10, r * 3
        for col, f in cols.items():
            if col in blank:
                continue
            ws[f"{col}{r}"] = "=" + (override or {}).get(col, f).format(r=r)
    return wb


def cache_calc11(path, cols=CALC_COLS, skip=()):
    vals = {"D": lambda a, b: a * 2, "F": lambda a, b: a + b,
            "H": lambda a, b: max(a, b), "J": lambda a, b: a - b,
            "L": lambda a, b: a * b}
    for r in range(3, 10):
        for col in cols:
            if col not in skip:
                inject_cache11(path, f"{col}{r}", vals[col](r * 10, r * 3))


p13 = save(wb_calc11(cols={}), "13문제.xlsx")
a13 = save(wb_calc11(), "13정답.xlsx")
cache_calc11(a13)
s13a = save(wb_calc11(blank=("D",)), "13학생a.xlsx")
cache_calc11(s13a, skip=("D",))
r, _, _ = grade(p13, a13, s13a)
c13 = r.wrong[0] if r.wrong else {}
check(13, "문제 1(D열) 전부 빈 셀 -> 32/40, 상태 '(빈 셀)' + 저장 확인 안내",
      r.earned == 32 and c13.get("label") == "계산 문제 1"
      and c13["cells"][0]["got"] == g.STATE_BLANK
      and "저장(Ctrl+S)한 파일이 맞는지" in (c13.get("note") or ""),
      str(c13.get("note"))[:160])
check(13, "카드에 '인식된 문제 5개 중 1번 (문제당 8점)' 표기",
      "인식된 문제 5개 중 1번 (문제당 8점)" in c13.get("note", ""),
      c13.get("note", "")[:120])
s13b = save(wb_calc11(override={"D": "2*A{r}"}), "13학생b.xlsx")
cache_calc11(s13b, skip=("D",))
# v2.0.4 전(재계산 없음) 경로: 계산값이 없으면 수식 문자열로만 비교 -> 감점
_re_on = g.RECALC_ENABLED
g.RECALC_ENABLED = False
try:
    r, _, _ = grade(p13, a13, s13b)
finally:
    g.RECALC_ENABLED = _re_on
c13b = r.wrong[0] if r.wrong else {}
check(13, "재계산 불가 시: 계산값 없음 상태 표기 + 학생 수식 전문 + 재저장 안내",
      r.earned == 32 and c13b["cells"][0]["got"] == g.STATE_NOCACHE
      and c13b["cells"][0]["got_formula"] == "=2*A3"
      and c13b["cells"][0].get("got_state") == "nocache"
      and "Excel에서 열어 저장" in c13b.get("note", ""),
      str(c13b.get("cells"))[:200])
check(13, "재계산 불가 시: 리포트에 '수식 문자열로 비교' 명시",
      any("수식 문자열로 비교" in n for n in r.notes), str(r.notes)[:200])
card13 = g._wrong_card_html(r.wrong[0])
cells13 = r.wrong[0]["cells"]
check(13, "리포트 카드: 내 답에 셀 상태 + 학생 수식 전문이 함께 보인다",
      cells13[0]["coord"] == "D3" and cells13[0]["got"] == g.STATE_NOCACHE
      and cells13[0]["got_formula"] == "=2*A3"
      and g.STATE_NOCACHE in card13 and "=2*A3" in card13
      and [c["coord"] for c in cells13[:3]] == ["D3", "D4", "D5"]
      and r.wrong[0].get("more"),
      str([c["coord"] for c in cells13]) + " more=" + str(r.wrong[0].get("more")))
s13ok = save(wb_calc11(override={"D": "2*A{r}"}), "13학생ok.xlsx")
cache_calc11(s13ok)
r, _, _ = grade(p13, a13, s13ok)
check(13, "같은 수식(=2*A3)에 계산값이 있으면 값 기준 정답 -> 40/40",
      r.earned == r.alloc, str(r.details))
# 문제 파일 오염: 문제 파일에 정답 수식(F/H/J/L)이 이미 들어 있음
p13x = save(wb_calc11(blank=("D",)), "13문제오염.xlsx")
cache_calc11(p13x, skip=("D",))
r, _, _ = grade(p13x, a13, s13a)
c13x = r.wrong[0] if r.wrong else {}
check(13, "오염된 문제 파일 -> 문제 1개로 인식(-40) + 오염 경고 노트",
      r.earned == 0 and abs(c13x.get("lost", 0) - 40) < 0.05
      and any("문제 파일 오염" in n for n in r.notes)
      and "인식된 문제 1개 중 1번 (문제당 40점)" in c13x.get("note", "")
      and "표준 5문제와 달라" in c13x.get("note", ""),
      " | ".join(r.notes)[:200])

# --- ④ 기본작업-2: 빈 셀 맞춤 잔여 서식 제외 (2024 상시 2회 K4:K9) -------------
VC = Alignment(vertical="center")
HC = Alignment(horizontal="center", vertical="center")


def wb_b2_align(table_align=None, blank_k=None, title_cc_cols=()):
    wb = wb_new("기본작업-2")
    ws = wb.active
    ws["B2"] = "대여 현황"
    for r in range(4, 7):
        for c in range(2, 6):
            ws.cell(r, c).value = f"d{r}{c}"
            if table_align is not None:
                ws.cell(r, c).alignment = table_align
    if blank_k is not None:
        for r in range(4, 7):
            ws.cell(r, 11).alignment = blank_k   # K4:K6 값 없는 빈 셀
    for c in title_cc_cols:
        ws.cell(2, c).alignment = Alignment(horizontal="centerContinuous")
    return wb


p14 = save(wb_b2_align(), "14문제.xlsx")
a14 = save(wb_b2_align(HC, blank_k=VC), "14정답.xlsx")
s14 = save(wb_b2_align(HC), "14학생.xlsx")
r, _, _ = grade(p14, a14, s14)
check(14, "정답 K4:K6(빈 셀) 세로 맞춤 잔여 서식은 채점 제외 -> 만점",
      r.earned == r.alloc, str(r.details))
s14b = save(wb_b2_align(), "14학생b.xlsx")
r, _, _ = grade(p14, a14, s14b)
check(14, "표(값 있는 셀) 맞춤을 안 하면 감점 유지",
      r.earned < r.alloc and "맞춤" in str(r.wrong), str(r.details)[:120])
a14c = save(wb_b2_align(HC, title_cc_cols=range(2, 7)), "14정답c.xlsx")
s14c = save(wb_b2_align(HC, title_cc_cols=(2,)), "14학생c.xlsx")
r, _, _ = grade(p14, a14c, s14c)
check(14, "'선택 영역의 가운데로'는 빈 셀(C2:F2)도 화면에 보이므로 감점 유지",
      r.earned < r.alloc and "C2:F2" in str(r.wrong), str(r.details)[:160])

# ---------------------------------------------------------------------------
# 12. v2.0.2 이의제기 5건 (2026 상시 복원 1회 — 공식 문제지 PDF 기준)
# ---------------------------------------------------------------------------
print("12. v2.0.2 이의제기 반영 (기대값 exclude/chart · 정의된 이름 표기 · cellIs "
      "연산자 · 계산작업 자료 셀/오염 · 차트 축 역할 비교)")
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.workbook.defined_name import DefinedName


# --- ① 기본작업-2: 정답 파일 잔여 서식(F2 가운데)을 세트 기대값 exclude로 제외 ---
def wb_b2x(f2_center=False, table=True):
    wb = wb_new("기본작업-2")
    ws = wb.active
    ws["A1"] = "택배 주문 현황"
    ws["F2"], ws["G2"] = "세액", 1.1
    for r in range(3, 6):
        for c in range(1, 8):
            ws.cell(r, c).value = f"v{r}{c}"
            if table:
                ws.cell(r, c).alignment = HC
    if f2_center:
        ws["F2"].alignment = Alignment(horizontal="center")
    return wb


p15 = save(wb_b2x(table=False), "15문제.xlsx")
a15 = save(wb_b2x(f2_center=True), "15정답.xlsx")
s15 = save(wb_b2x(), "15학생.xlsx")
r, _, _ = grade(p15, a15, s15)
check(15, "지시 없는 F2 가운데 맞춤(정답 잔여 서식) — 기대값 없으면 감점 (재현)",
      r.earned < r.alloc and "F2" in str(r.details), str(r.details)[:120])
KEY15 = {"exclude": {"기본작업-2": [{"range": "F2", "kinds": ["alignment"],
                                   "why": "문제지 지시 없음"}]}}
r, _, _ = grade(p15, a15, s15, KEY15)
check(15, "기대값 exclude(F2 맞춤) -> 만점 + 제외 안내 노트",
      r.earned == r.alloc and any("채점 제외: F2 맞춤" in n and "문제지 지시 없음" in n
                                  for n in r.notes),
      " | ".join(r.notes)[:160])
s15b = save(wb_b2x(table=False), "15학생b.xlsx")
r, _, _ = grade(p15, a15, s15b, KEY15)
check(15, "exclude는 F2 맞춤만 — 표 맞춤을 안 하면 감점 유지(F2는 인용 안 함)",
      r.earned < r.alloc and "F2" not in str(r.details), str(r.details)[:120])
check(15, "normalize_key: 문자열 항목·kinds 생략 허용, chart 섹션 보존",
      g.normalize_key({"exclude": {"기본작업-2": ["F2"]},
                       "chart": {"차트작업": {"axes": []}}})["exclude"]
      == {"기본작업-2": [{"range": "F2", "kinds": None, "why": None}]}
      and "차트작업" in g.normalize_key({"chart": {"차트작업": {}}})["chart"])


# --- ② 기본작업-2: 정의된 이름 — 내 답에 실제 정의 표기 + 영역 수 차이 설명 ---
def wb_names(ref=None):
    wb = wb_new("기본작업-2")
    ws = wb.active
    for r in range(4, 11):
        ws.cell(r, 1).value = f"ORD-{r}"
        ws.cell(r, 3).value = f"상품{r}"
    if ref:
        wb.defined_names["주문상품"] = DefinedName("주문상품", attr_text=ref)
    return wb


REF16 = "'기본작업-2'!$A$4:$A$10,'기본작업-2'!$C$4:$C$10"
p16 = save(wb_names(), "16문제.xlsx")
a16 = save(wb_names(REF16), "16정답.xlsx")
s16 = save(wb_names("'기본작업-2'!$C$4:$C$10"), "16학생.xlsx")
r, _, _ = grade(p16, a16, s16)
c16 = r.wrong[0] if r.wrong else {}
check(16, "한 영역만 정의 -> 감점 + 내 답 \"'주문상품' = 기본작업-2!C4:C10\" 표기",
      r.earned < r.alloc and bool(c16.get("props"))
      and c16["props"][0]["got"] == "'주문상품' = 기본작업-2!C4:C10"
      and c16["props"][0]["expected"]
      == "'주문상품' = 기본작업-2!A4:A10,기본작업-2!C4:C10",
      str(c16.get("props")))
check(16, "'영역 2개 중 1개' 차이 설명 + Ctrl 다중 선택 안내",
      any("영역 2개" in d and "1개" in d for d in c16.get("diff_notes") or [])
      and "Ctrl" in (c16.get("hint") or ""), str(c16.get("diff_notes")))
s16b = save(wb_names(), "16학생b.xlsx")
r, _, _ = grade(p16, a16, s16b)
c16b = r.wrong[0] if r.wrong else {}
check(16, "이름 없음 -> 내 답 \"'주문상품' 이름 없음\"",
      bool(c16b.get("props")) and c16b["props"][0]["got"] == "'주문상품' 이름 없음",
      str(c16b.get("props")))
s16c = save(wb_names(REF16), "16학생c.xlsx")
r, _, _ = grade(p16, a16, s16c)
check(16, "두 영역 모두 정의(A4:A10,C4:C10) -> 만점", r.earned == r.alloc,
      str(r.details))


# --- ③ 기본작업-3: 셀 값 기준 규칙 연산자 표기·판정 ('70% 이상' = >=) ---
E_VALS = [0.15, 0.23, 0.45, 0.63, 0.66, 0.69, 0.78, 0.81]
BLUE = PatternFill("solid", fgColor="FF0070C0")
YELLOW = PatternFill("solid", fgColor="FFFFFF00")
WHITE_B = Font(bold=True, color="FFFFFFFF")


def wb_cf3(op_hi=None, vals=E_VALS):
    wb = wb_new("기본작업-3")
    ws = wb.active
    ws["E3"] = "공정률(%)"
    for i, v in enumerate(vals):
        ws.cell(4 + i, 5).value = v
    if op_hi:
        ws.conditional_formatting.add(
            "E4:E11", CellIsRule(operator=op_hi, formula=["0.7"], fill=BLUE,
                                 font=WHITE_B))
        ws.conditional_formatting.add(
            "E4:E11", CellIsRule(operator="lessThan", formula=["0.3"],
                                 fill=YELLOW, font=Font(bold=True)))
    return wb


p17 = save(wb_cf3(), "17문제.xlsx")
a17 = save(wb_cf3("greaterThanOrEqual"), "17정답.xlsx")
s17 = save(wb_cf3("greaterThan"), "17학생.xlsx")
r, _, _ = grade(p17, a17, s17)
c17 = r.wrong[0] if r.wrong else {}
check(17, "'70% 이상' 규칙에 > 0.7 -> 2/5 + 조건 '셀 값 >= 0.7' / 내 답 '셀 값 > 0.7' 표기",
      r.earned == 2 and "셀 값 >= 0.7" in str(r.details)
      and "셀 값 > 0.7" in str(r.details)
      and any(p["name"] == "조건" and p["expected"] == "셀 값 >= 0.7"
              and p["got"] == "셀 값 > 0.7" for p in c17.get("props", [])),
      str(r.details)[:200])
d17 = c17.get("diff_notes") or []
check(17, "'비교 연산자 불일치'(>= 이상 / > 초과) 진단 — '규칙 유형 불일치' 아님, "
          "다른 규칙(< 0.3)은 인용 안 함",
      any(d.startswith("비교 연산자 불일치") and ">= (이상)" in d and "> (초과)" in d
          and "< (미만)" not in d for d in d17)
      and not any("규칙 유형 불일치" in d for d in d17)
      and any(p["name"] == "비교 연산자" and p["got"] == "> (초과)"
              for p in c17.get("props", [])),
      str(d17))
check(17, "경계값 0.7과 같은 셀이 없어 화면은 같지만 규칙 자체를 채점한다는 안내",
      any("경계값 0.7과 정확히 같은 셀이 없어" in d for d in d17), str(d17)[-160:])
card17 = g._wrong_card_html(r.wrong[0])
check(17, "리포트 카드·풀이 단계에 연산자 포함 조건 표기",
      any(p.get("name") == "조건" and "셀 값 >= 0.7" in str(p.get("expected"))
          for p in r.wrong[0].get("props") or [])
      and "셀 값 &gt;= 0.7" in card17
      and any("셀 값 >= 0.7" in s for s in c17.get("explain") or []),
      str(r.wrong[0].get("props"))[:200])
vals17b = list(E_VALS)
vals17b[5] = 0.7
p17b = save(wb_cf3(vals=vals17b), "17문제b.xlsx")
a17b = save(wb_cf3("greaterThanOrEqual", vals17b), "17정답b.xlsx")
s17b = save(wb_cf3("greaterThan", vals17b), "17학생b.xlsx")
r, _, _ = grade(p17b, a17b, s17b)
d17b = (r.wrong[0].get("diff_notes") if r.wrong else None) or []
check(17, "경계값 0.7 셀(E9)이 있으면 화면 결과도 다르다고 안내",
      any("경계값 0.7과 같은 셀(E9)" in d for d in d17b), str(d17b)[-160:])
s17c = save(wb_cf3("greaterThanOrEqual"), "17학생c.xlsx")
r, _, _ = grade(p17, a17, s17c)
check(17, ">= 0.7 그대로 -> 만점", r.earned == r.alloc, str(r.details))


# --- ④ 계산작업: 정답이 수식이 아닌 셀(자료 차이)은 문제로 묶지 않음 + 오염(다른 수식) 경고 ---
wb18 = wb_calc11(cols={})
wb18.active["M5"] = 0                       # 문제 파일: 자료 셀 0
p18 = save(wb18, "18문제.xlsx")
wb18 = wb_calc11()
wb18.active["M5"] = "미응시"                 # 정답 파일: 자료 셀만 다름 (수식 아님)
a18 = save(wb18, "18정답.xlsx")
cache_calc11(a18)
wb18 = wb_calc11()
wb18.active["M5"] = 0
s18 = save(wb18, "18학생.xlsx")
cache_calc11(s18)
r, _, _ = grade(p18, a18, s18)
check(18, "정답 M5 '미응시'(자료 셀) vs 내 답 0 — 수식 5문제 정답이면 40/40 + 자료 셀 제외 노트",
      r.earned == r.alloc and any("정답이 수식이 아닌 셀 M5" in n for n in r.notes)
      and any("5개 그룹으로 인식" in n for n in r.notes), " | ".join(r.notes)[:200])
p18x = save(wb_calc11(cols={"D": "A{r}*3"}), "18문제오염.xlsx")   # 정답과 다른 수식
a18x = save(wb_calc11(), "18정답x.xlsx")
cache_calc11(a18x)
s18x = save(wb_calc11(), "18학생x.xlsx")
cache_calc11(s18x)
res18, sc18, gn18, _ = g.run_grading(p18x, a18x, s18x, g.normalize_key({}))
r = res18[0]
check(18, "문제 파일에 정답과 다른 수식(D3:D9)이 있으면 5문제여도 오염 경고 + 플래그",
      r.earned == r.alloc
      and any("정답과 다른 수식이 이미 7개" in n and "문제 파일 오염" in n for n in r.notes)
      and getattr(r, "contaminated", None) == ["D3:D9"]
      and not any("정답과 같은 수식" in n for n in r.notes), " | ".join(r.notes)[:200])
html18 = os.path.join(TMP, "18.html")
g.write_html(html18, res18, sc18, gn18, (p18x, a18x, s18x))
check(18, "리포트 상단에 '문제 파일 오염 의심' 배너",
      "문제 파일 오염 의심" in open(html18, encoding="utf-8").read())


# --- ⑤ 차트 축: 기본/보조 역할 짝 비교, 정답 '자동' 속성 제외, 최소 0 = 자동 0, 기대값 chart ---
def wb_combo(primary_max=None, sec=None, narrow=False):
    """막대(기본 축) + 선택적 꺾은선(보조 축). sec=(최대, 최소, 주 단위)."""
    wb = wb_new("차트작업")
    ws = wb.active
    rows = [("베트남", 22650000, 7000), ("미국", 16800000, 10000),
            ("이탈리아", 26500000, 6000), ("스위스", 23050000, 8000)]
    if narrow:
        rows = [(n, a, b2) for (n, a, _b), b2
                in zip(rows, (9000, 9500, 9800, 10000))]
    for i, (lab, v1, v2) in enumerate(rows, start=2):
        ws.cell(i, 1).value = lab
        ws.cell(i, 2).value = v1
        ws.cell(i, 3).value = v2
    bar = BarChart()
    bar.add_data(Reference(ws, min_col=2, min_row=2, max_row=5))
    bar.set_categories(Reference(ws, min_col=1, min_row=2, max_row=5))
    if primary_max is not None:
        bar.y_axis.scaling.max = primary_max
    if sec is not None:
        line = LineChart()
        line.add_data(Reference(ws, min_col=3, min_row=2, max_row=5))
        line.y_axis.axId = 200
        line.y_axis.crosses = "max"
        mx, mn, un = sec
        if mx is not None:
            line.y_axis.scaling.max = mx
        if mn is not None:
            line.y_axis.scaling.min = mn
        if un is not None:
            line.y_axis.majorUnit = un
        bar += line
    ws.add_chart(bar, "E5")
    return wb


p19 = save(wb_combo(), "19문제.xlsx")
a19 = save(wb_combo(sec=(12000, 0, 1500)), "19정답.xlsx")
s19 = save(wb_combo(primary_max=40000000, sec=(12000, None, 1500)), "19학생.xlsx")
r, _, _ = grade(p19, a19, s19)
check(19, "보조 축 최대 12,000·단위 1,500 수행, 최소 자동(=0)·기본 축 최대 40 부수 설정 -> 만점",
      r.earned == r.alloc, str(r.details))
s19b = save(wb_combo(sec=(10000, None, 1500)), "19학생b.xlsx")
r, _, _ = grade(p19, a19, s19b)
check(19, "보조 축 최대 10,000 -> 감점, '보조 축 최대 12,000 / 10,000' 표기(4e+07 표기 없음)",
      r.earned < r.alloc and "보조 축 최대 12,000" in str(r.details)
      and "보조 축 최대 10,000" in str(r.details) and "e+0" not in str(r.details),
      str(r.details)[:160])
s19c = save(wb_combo(sec=(12000, None, 1000)), "19학생c.xlsx")
r, _, _ = grade(p19, a19, s19c)
check(19, "주 단위 1,000 -> 감점 ('보조 축 주 단위 1,500' 표기)",
      r.earned < r.alloc and "보조 축 주 단위 1,500" in str(r.details),
      str(r.details)[:160])
p19x = save(wb_combo(primary_max=40000000, sec=(12000, None, 1500)), "19문제오염.xlsx")
r, _, _ = grade(p19x, a19, s19)
check(19, "오염 문제 파일(기본 축 최대 40 이미 있음): '최대 자동' 지시로 오인하지 않음 -> 만점",
      r.earned == r.alloc, str(r.details))
p19n = save(wb_combo(narrow=True), "19문제n.xlsx")
a19n = save(wb_combo(sec=(12000, 0, 1500), narrow=True), "19정답n.xlsx")
s19n = save(wb_combo(sec=(12000, None, 1500), narrow=True), "19학생n.xlsx")
r, _, _ = grade(p19n, a19n, s19n)
check(19, "값이 좁게 몰린 계열(9,000~10,000)은 자동 최소가 0이 아니므로 '최소 0' 감점 유지",
      r.earned < r.alloc and "보조 축 최소 0" in str(r.details), str(r.details)[:160])
KEY19 = {"chart": {"차트작업": {"axes": [{"axis": "secondary", "max": 12000,
                                         "major_unit": 1500}]}}}
r, _, _ = grade(p19x, a19, s19, KEY19)
check(19, "기대값 chart.axes(보조 축 최대·단위)만 채점 -> 만점", r.earned == r.alloc,
      str(r.details))
r, _, _ = grade(p19x, a19, s19c, KEY19)
check(19, "기대값 명시 단위 1,500 vs 1,000 -> 감점",
      r.earned < r.alloc and "보조 축 주 단위 1,500" in str(r.details)
      and "보조 축 주 단위 1,000" in str(r.details), str(r.details)[:160])
s19d = save(wb_combo(), "19학생d.xlsx")
r, _, _ = grade(p19x, a19, s19d, KEY19)
check(19, "보조 축이 없으면 '(축 없음)' 표기", "(축 없음)" in str(r.details),
      str(r.details)[:160])

# ---------------------------------------------------------------------------
# 13. v2.0.3 이의제기 반영 (코코 모의고사 1회 — 표준색 파랑·표시 형식 동치·
#     부등호 어순 동치·매크로 VBA 코드 인정과 표기)
# ---------------------------------------------------------------------------
print("13. v2.0.3 이의제기 반영 (표시 형식 동치 · 부등호 어순 · 매크로 VBA 코드)")
import re

# --- 13-1 표준색 파랑(#0070C0) / '병합하고 가운데 맞춤'은 가로만 ---
check(20, "Excel [표준 색] 파랑 #0070C0 -> '파랑', [다른 색] #0000FF는 이름 없음",
      "파랑" == g.color_name_ko("0070C0").split("(")[0]
      and g.color_name_ko("0000FF") == "#0000FF",
      f'{g.color_name_ko("0070C0")} / {g.color_name_ko("0000FF")}')


def wb_title(color, vertical, nf, values=None, nf_cells=True):
    wb = wb_new("기본작업-2")
    ws = wb.active
    ws["A1"] = "코코문구 3월 판매 현황"
    for j, h in enumerate(["상품코드", "판매량"]):
        ws.cell(3, 1 + j).value = h
    for i, v in enumerate(values or [460, 380, 310, 270]):
        ws.cell(4 + i, 1).value = f"ST-0{i + 1}"
        ws.cell(4 + i, 2).value = v
    if color:
        ws.merge_cells("A1:B1")
        ws["A1"].font = Font(name="돋움", size=16, bold=True, color=color)
        ws["A1"].alignment = Alignment(horizontal="center",
                                       vertical=vertical) if vertical \
            else Alignment(horizontal="center")
    if nf and nf_cells:
        for i in range(len(values or [0] * 4)):
            ws.cell(4 + i, 2).number_format = nf
    return wb


p20 = save(wb_title(None, None, None), "20문제.xlsx")
a20 = save(wb_title("FF0070C0", None, '0"개"'), "20정답.xlsx")
s20 = save(wb_title("FF0070C0", None, '0"개"'), "20학생.xlsx")
r, _, _ = grade(p20, a20, s20)
check(20, "정답 파일에 문제지에 없는 세로 맞춤이 없으면 가로 가운데만 한 풀이도 만점",
      r.earned == r.alloc, str(r.details))
a20v = save(wb_title("FF0070C0", "center", '0"개"'), "20정답세로.xlsx")
r, _, _ = grade(p20, a20v, s20)
check(20, "정답 파일에 세로 맞춤 잔여 서식이 있으면 감점(정답 파일을 고쳐야 하는 "
          "상태를 재현)",
      r.earned < r.alloc and "세로 맞춤" in str([c.get("props")
                                             for c in r.wrong]),
      str(r.details))

# --- 13-2 표시 형식: 코드는 달라도 화면 표시가 같으면 정답 ---
check(21, "G/표준 은 General 과 같은 코드로 정규화",
      g._norm_number_format('G/표준"개"') == g._norm_number_format('General"개"'),
      g._norm_number_format('G/표준"개"'))
check(21, "리포트 표기는 원문 코드 그대로, General 은 한국어 G/표준",
      g.nf_display('0"개"') == '0"개"'
      and g.nf_display('General"개"') == 'G/표준"개"',
      g.nf_display('0"개"') + " / " + g.nf_display('General"개"'))
check(21, "렌더 동치: 정수 460은 0\"개\" 와 G/표준\"개\" 모두 '460개'",
      g.nf_render(460, '0"개"') == "460개"
      and g.nf_same_display(460, '0"개"', 'General"개"')
      and g.nf_same_display(0, '0"개"', 'General"개"'),
      str(g.nf_render(460, 'General"개"')))
check(21, "렌더 동치 아님: 460.5는 '461개' vs '460.5개'",
      not g.nf_same_display(460.5, '0"개"', 'General"개"'))
check(21, "천 단위 배율·백분율 렌더",
      g.nf_render(3650000, '#,##0,"천원"') == "3,650천원"
      and g.nf_render(0.5714, "0.00%") == "57.14%"
      and not g.nf_same_display(3650000, '#,##0,"천원"', 'General"천원"'),
      g.nf_render(3650000, '#,##0,"천원"'))

s21 = save(wb_title("FF0070C0", None, 'General"개"'), "21학생.xlsx")
r, _, notes = grade(p20, a20, s21)
check(21, 'E열 정수에 G/표준"개"(저장형 General"개") -> 만점 + 동치 안내',
      r.earned == r.alloc
      and any("화면 표시가 같아 정답" in n for n in r.notes),
      str(r.notes))
p20f = save(wb_title(None, None, None, values=[460, 380.5, 310, 270]),
            "21문제f.xlsx")
a20f = save(wb_title("FF0070C0", None, '0"개"', values=[460, 380.5, 310, 270]),
            "21정답f.xlsx")
s21f = save(wb_title("FF0070C0", None, 'General"개"',
                     values=[460, 380.5, 310, 270]), "21학생f.xlsx")
r, _, _ = grade(p20f, a20f, s21f)
nf_props = [p for c in r.wrong for p in (c.get("props") or [])
            if p["name"].startswith("표시 형식")]
check(21, "소수 값이 있으면 표시가 달라 감점 유지",
      r.earned < r.alloc and nf_props, str(r.details))
check(21, '카드에 원문 코드 0"개" / G/표준"개" 표기 (따옴표 없는 0개·General개 아님)',
      nf_props and '0"개"' in nf_props[0]["expected"]
      and 'G/표준"개"' in nf_props[0]["got"]
      and "General" not in nf_props[0]["got"],
      str(nf_props[:1]))
check(21, "형식 코드 해설도 원문 기준 (문자 '개' 붙여 표시)",
      "문자 '개'" in nf_props[0]["expected"], nf_props[0]["expected"])

# --- 13-3 조건부 서식: 부등호 좌우 교환 동치, $ 고정은 여전히 오답 ---
check(22, "비교 연산자 좌우 교환 동치 (>=↔<=, >↔<, =, <>)",
      g.norm_formula("=$F3>=AVERAGE($F$3:$F$12)")
      == g.norm_formula("=AVERAGE($F$3:$F$12)<=F3")
      and g.norm_formula("=IF(A1>=90,1,0)") == g.norm_formula("=IF(90<=A1,1,0)")
      and g.norm_formula("=A1<>B1") == g.norm_formula("=B1<>A1"),
      g.norm_formula("=$F3>=AVERAGE($F$3:$F$12)"))
check(22, "좌우 교환이 아닌 차이는 여전히 다른 수식 (A1>B1 vs B1>A1)",
      g.norm_formula("=A1>B1") != g.norm_formula("=B1>A1"))


def wb_avg_cf(formula=None, color="FF0070C0"):
    wb = wb_new("기본작업-3")
    ws = wb.active
    ws["A1"] = "코코정보고 1학년 모의고사 성적"
    for j, h in enumerate(["반", "이름", "총점"]):
        ws.cell(2, 1 + j).value = h
    for i, (cl, nm, tot) in enumerate([("1반", "김하늘", 249),
                                       ("1반", "나준서", 217),
                                       ("2반", "도예린", 273),
                                       ("2반", "류시원", 174)]):
        ws.cell(3 + i, 1).value = cl
        ws.cell(3 + i, 2).value = nm
        ws.cell(3 + i, 3).value = tot
    if formula:
        from openpyxl.formatting.rule import FormulaRule
        ws.conditional_formatting.add(
            "A3:C6", FormulaRule(formula=[formula],
                                 font=Font(bold=True, color=color)))
    return wb


p22 = save(wb_avg_cf(), "22문제.xlsx")
a22 = save(wb_avg_cf("$C3>=AVERAGE($C$3:$C$6)"), "22정답.xlsx")
s22 = save(wb_avg_cf("AVERAGE($C$3:$C$6)<=$C3"), "22학생.xlsx")
r, _, _ = grade(p22, a22, s22)
check(22, "부등호 어순만 반대인 조건부 서식 -> 만점", r.earned == r.alloc,
      str(r.details))
s22b = save(wb_avg_cf("AVERAGE($C$3:$C$6)<=C3"), "22학생b.xlsx")
r, _, _ = grade(p22, a22, s22b)
notes22 = str([c.get("diff_notes") for c in r.wrong])
check(22, "상대참조 C3(열 고정 없음)은 오답 유지 + '참조 고정($)' 진단",
      r.earned < r.alloc and "참조 고정($)" in notes22, notes22[:160])
check(22, "'부등호 좌우 순서만 바꾼 식은 인정, 남은 차이는 $' 안내",
      "부등호의 좌우 순서만" in notes22, notes22[:200])
point22 = str([c.get("point") for c in r.wrong])
check(22, "해설이 이 문제의 실제 범위·기준 열로 생성 (다른 세트 예시 LEFT($A4) 인용 안 함)",
      "A3:C6" in point22 and "$C3" in point22 and "LEFT($A4" not in point22
      and "AVERAGE($C$3:$C$6)" in point22, point22[:220])

# --- 13-4 매크로: VBA 코드 인정 + 근거 표기 ---
CFB_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def ovba_container(text):
    """소스 -> MS-OVBA CompressedContainer (비압축 청크)."""
    import struct
    data = text.encode("cp949")
    out = bytearray(b"\x01")
    for i in range(0, len(data), 4096):
        chunk = data[i:i + 4096]
        out += struct.pack("<H", (len(chunk) - 1) | 0x3000) + chunk
    return bytes(out)


def _dir_entry(name, etype, start, size, child=0xFFFFFFFF):
    import struct
    nb = name.encode("utf-16-le") + b"\x00\x00"
    e = bytearray(128)
    e[0:len(nb)] = nb
    struct.pack_into("<H", e, 64, len(nb))
    e[66] = etype
    e[67] = 1
    struct.pack_into("<III", e, 68, 0xFFFFFFFF, 0xFFFFFFFF, child)
    struct.pack_into("<I", e, 116, start)
    struct.pack_into("<Q", e, 120, size)
    return bytes(e)


def build_vba_bin_mini(source, module="Module1", ssz=512, msz=64):
    """모듈 스트림을 **미니 스트림**에 담은 합성 vbaProject.bin.

    실제 Excel 이 저장한 vbaProject.bin 은 4096바이트(미니 스트림 컷오프)보다
    작은 스트림을 64바이트 미니 섹터에 담는다 — 매크로 모듈은 대개 그보다
    작으므로 실물은 거의 항상 이 경로를 탄다. 지금까지 합성 픽스처는 늘
    일반 FAT 체인만 만들어서 이 경로가 한 번도 검증되지 않았다.

    배치: [헤더] [섹터0=FAT] [섹터1=디렉터리] [섹터2=MiniFAT] [섹터3~=미니 스트림]
    """
    import struct
    shift = ssz.bit_length() - 1
    mshift = msz.bit_length() - 1
    per = ssz // 4
    body = ovba_container(source)
    mini = body + b"\x00" * ((-len(body)) % msz)
    n_mini = len(mini) // msz
    mini_pad = mini + b"\x00" * ((-len(mini)) % ssz)
    n_cont = max(1, len(mini_pad) // ssz)
    # FAT: 0=FAT 자신, 1=디렉터리, 2=MiniFAT, 3..=미니 스트림 컨테이너
    fat = [0xFFFFFFFD, 0xFFFFFFFE, 0xFFFFFFFE]
    for i in range(n_cont):
        fat.append(0xFFFFFFFE if i == n_cont - 1 else 3 + i + 1)
    fat += [0xFFFFFFFF] * (per - len(fat))
    fat_sector = b"".join(struct.pack("<I", v) for v in fat)
    minifat = [i + 1 for i in range(n_mini - 1)] + [0xFFFFFFFE]
    minifat += [0xFFFFFFFF] * (per - len(minifat))
    minifat_sector = b"".join(struct.pack("<I", v) for v in minifat)
    d = bytearray()
    d += _dir_entry("Root Entry", 5, 3, len(mini), child=1)
    d += _dir_entry("VBA", 1, 0xFFFFFFFE, 0, child=2)
    d += _dir_entry(module, 2, 0, len(body))      # 미니 섹터 0번부터
    d += b"\x00" * (ssz - len(d))
    h = bytearray(512)
    h[0:8] = CFB_MAGIC
    struct.pack_into("<H", h, 26, 3 if ssz == 512 else 4)
    struct.pack_into("<H", h, 28, 0xFFFE)
    struct.pack_into("<H", h, 30, shift)
    struct.pack_into("<H", h, 32, mshift)
    struct.pack_into("<I", h, 44, 1)        # FAT 섹터 수
    struct.pack_into("<I", h, 48, 1)        # 디렉터리 첫 섹터
    struct.pack_into("<I", h, 56, 4096)     # 미니 스트림 컷오프
    struct.pack_into("<I", h, 60, 2)        # MiniFAT 첫 섹터
    struct.pack_into("<I", h, 64, 1)        # MiniFAT 섹터 수
    struct.pack_into("<I", h, 68, 0xFFFFFFFE)
    struct.pack_into("<I", h, 76, 0)
    for i in range(1, 109):
        struct.pack_into("<I", h, 76 + 4 * i, 0xFFFFFFFF)
    head = bytes(h) + b"\x00" * (ssz - 512)
    return head + fat_sector + bytes(d) + minifat_sector + mini_pad


def build_vba_bin(source, module="Module1", ssz=512):
    """모듈 1개를 담은 합성 vbaProject.bin (OLE/CFB).

    ssz: 섹터 크기. 512(CFB v3)와 4096(CFB v4) 모두 만들 수 있다 — 실제
    Excel이 내보낸 vbaProject.bin 중에 4096바이트 섹터 파일이 있고,
    채점기가 섹터 시작 위치를 512로 못박아 두면 그런 파일에서 스트림을
    하나도 읽지 못한다(2.0.6에서 고친 버그).
    """
    import struct
    shift = ssz.bit_length() - 1
    per = ssz // 4
    body = ovba_container(source)
    stream = body + b"\x00" * ((-len(body)) % ssz)
    n = len(stream) // ssz
    fat = [0xFFFFFFFD, 0xFFFFFFFE] + [i + 3 for i in range(n - 1)] \
        + [0xFFFFFFFE]
    fat += [0xFFFFFFFF] * (per - len(fat))
    fat_sector = b"".join(struct.pack("<I", v) for v in fat)
    d = bytearray()
    d += _dir_entry("Root Entry", 5, 0xFFFFFFFE, 0, child=1)
    d += _dir_entry("VBA", 1, 0xFFFFFFFE, 0, child=2)
    d += _dir_entry(module, 2, 2, len(body))
    d += b"\x00" * (ssz - len(d))
    h = bytearray(512)
    h[0:8] = CFB_MAGIC
    struct.pack_into("<H", h, 26, 3 if ssz == 512 else 4)
    struct.pack_into("<H", h, 28, 0xFFFE)
    struct.pack_into("<H", h, 30, shift)
    struct.pack_into("<H", h, 32, 6)
    struct.pack_into("<I", h, 44, 1)
    struct.pack_into("<I", h, 48, 1)
    struct.pack_into("<I", h, 56, 4096)
    struct.pack_into("<I", h, 60, 0xFFFFFFFE)
    struct.pack_into("<I", h, 68, 0xFFFFFFFE)
    struct.pack_into("<I", h, 76, 0)
    for i in range(1, 109):
        struct.pack_into("<I", h, 76 + 4 * i, 0xFFFFFFFF)
    # 헤더는 항상 512바이트지만 섹터 0은 '한 섹터' 뒤에서 시작한다.
    head = bytes(h) + b"\x00" * (ssz - 512)
    return head + fat_sector + bytes(d) + stream


MACRO_SRC = 'Attribute VB_Name = "Module1"\n' + "\n".join([
    "Sub 판매금액()",
    '    Range("D4").Select',
    '    ActiveCell.FormulaR1C1 = "=RC[-2]*RC[-1]"',
    '    Selection.AutoFill Destination:=Range("D4:D6"), Type:=xlFillDefault',
    "End Sub",
    "Sub 서식()",
    '    Range("A3:D3").Select',
    "    With Selection.Interior",
    "        .Pattern = xlSolid",
    "        .Color = 65535",
    "    End With",
    "    Selection.Font.Bold = True",
    "End Sub",
]) + "\n" + "\n".join(f"' 여백 {i} " + "-" * 40 for i in range(80))

VML = """<xml xmlns:v="urn:schemas-microsoft-com:vml"
 xmlns:o="urn:schemas-microsoft-com:office:office"
 xmlns:x="urn:schemas-microsoft-com:office:excel">
 <v:shape id="_x0000_s1025" type="#_x0000_t201" o:button="t">
  <v:textbox><div><font face="맑은 고딕">{text}</font></div></v:textbox>
  <x:ClientData ObjectType="Button"><x:Anchor>{anchor}</x:Anchor>
   <x:FmlaMacro>[0]!{macro}</x:FmlaMacro></x:ClientData></v:shape></xml>"""


def to_xlsm(src, out, vba_bin, sheet_part=None, button=None):
    """합성 xlsx -> xlsm (vbaProject.bin + 양식 컨트롤 단추 VML 주입)."""
    import zipfile
    with zipfile.ZipFile(src) as zin:
        blobs = {n: zin.read(n) for n in zin.namelist()}
    ct = blobs["[Content_Types].xml"].decode("utf-8")
    ct = ct.replace("</Types>",
                    '<Default Extension="bin" ContentType='
                    '"application/vnd.ms-office.vbaProject"/>'
                    '<Default Extension="vml" ContentType='
                    '"application/vnd.openxmlformats-officedocument.'
                    'vmlDrawing"/></Types>')
    blobs["[Content_Types].xml"] = ct.encode("utf-8")
    rl = blobs["xl/_rels/workbook.xml.rels"].decode("utf-8")
    blobs["xl/_rels/workbook.xml.rels"] = rl.replace(
        "</Relationships>",
        '<Relationship Id="rIdVba" Type="http://schemas.microsoft.com/office/'
        '2006/relationships/vbaProject" Target="vbaProject.bin"/>'
        "</Relationships>").encode("utf-8")
    blobs["xl/vbaProject.bin"] = vba_bin
    if button and sheet_part:
        text, macro, anchor = button
        blobs["xl/drawings/vmlDrawing1.vml"] = VML.format(
            text=text, macro=macro, anchor=anchor).encode("utf-8")
        srels = re.sub(r"worksheets/(sheet\d+\.xml)$",
                       r"worksheets/_rels/\1.rels", sheet_part)
        cur = blobs.get(srels, (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/'
            '2006/relationships"></Relationships>').encode("utf-8")
        ).decode("utf-8")
        blobs[srels] = cur.replace(
            "</Relationships>",
            '<Relationship Id="rIdVml" Type="http://schemas.openxmlformats.'
            'org/officeDocument/2006/relationships/vmlDrawing" '
            'Target="/xl/drawings/vmlDrawing1.vml"/></Relationships>'
        ).encode("utf-8")
        sx = blobs[sheet_part].decode("utf-8")
        if 'xmlns:r=' not in sx.split(">", 1)[0]:
            sx = sx.replace("<worksheet ", '<worksheet xmlns:r="http://'
                            'schemas.openxmlformats.org/officeDocument/2006/'
                            'relationships" ', 1)
        blobs[sheet_part] = sx.replace(
            "</worksheet>", '<legacyDrawing r:id="rIdVml"/></worksheet>'
        ).encode("utf-8")
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for n, b in blobs.items():
            zout.writestr(n, b)
    return out


VBA_BIN = build_vba_bin(MACRO_SRC)
mods = g.vba_module_sources(VBA_BIN)
check(23, "vbaProject.bin(OLE/CFB + MS-OVBA)에서 모듈 소스 복원",
      len(mods) == 1 and "Sub 서식()" in mods[0][1], str([m[0] for m in mods]))
vunits = g.vba_macro_units(mods[0][1]) if mods else []
check(23, "매크로 이름·대상 범위·동작 추출 (판매금액=값/수식, 서식=채우기·글꼴)",
      [u["name"] for u in vunits] == ["판매금액", "서식"]
      and g.vba_covers(vunits, "A3", "fill") == "서식"
      and g.vba_covers(vunits, "C3", "font") == "서식"
      and g.vba_covers(vunits, "D5", "value") == "판매금액"
      and g.vba_covers(vunits, "E3", "font") is None,
      str([(u["name"], u["ranges"], sorted(u["kinds"])) for u in vunits]))


def wb_macro(done=False):
    wb = wb_new("매크로작업")
    ws = wb.active
    ws["A1"] = "코코전자 대리점 판매 현황"
    for j, h in enumerate(["대리점", "판매량", "단가(천원)", "판매금액(천원)"]):
        ws.cell(3, 1 + j).value = h
    for i, (nm, q, pr) in enumerate([("강북점", 340, 620), ("강서점", 280, 750),
                                     ("분당점", 410, 680)]):
        ws.cell(4 + i, 1).value = nm
        ws.cell(4 + i, 2).value = q
        ws.cell(4 + i, 3).value = pr
    if done:
        for i in range(3):
            ws.cell(4 + i, 4).value = f"=B{4 + i}*C{4 + i}"
        for c in range(1, 5):
            cell = ws.cell(3, c)
            cell.fill = PatternFill("solid", fgColor="FFFF00")
            base = cell.font
            cell.font = Font(name=base.name, size=base.size, bold=True)
    return wb


p23 = save(wb_macro(False), "23문제.xlsx")
a23 = save(wb_macro(True), "23정답.xlsx")
s23 = save(wb_macro(True), "23학생.xlsx")
r, _, _ = grade(p23, a23, s23)
check(23, "정답 파일의 '굵게'가 기본 글꼴 이름을 유지하면 실행 결과가 같은 풀이는 만점",
      r.earned == r.alloc, str(r.details))

# 정답 파일이 글꼴 이름을 잃은 상태(수정 전 코코 1회) 재현
wb_bad = wb_macro(True)
for c in range(1, 5):
    wb_bad["매크로작업"].cell(3, c).font = Font(bold=True)
a23bad = save(wb_bad, "23정답글꼴손실.xlsx")
r, _, _ = grade(p23, a23bad, s23)
check(23, "정답 파일이 글꼴 이름을 잃으면 오탐 (정답 파일을 고쳐야 하는 상태를 재현)",
      r.earned < r.alloc, str(r.details))
card = next((c for c in r.wrong if c["label"] == "매크로 서식 결과"), None)
check(23, "매크로 서식 카드에 정답/내 답 서식 값을 실제로 표기 ('-' 금지)",
      card is not None
      and all(c.get("expected") and c.get("got") for c in card["cells"])
      and any("pt" in str(c["expected"]) for c in card["cells"]),
      str(card["cells"][:2]) if card else "카드 없음")

# 매크로만 기록하고 실행하지 않은 풀이 (결과 서식/값 없음) + 단추
s23x = save(wb_macro(False), "23학생코드만.xlsx")
_bs = g.Book(s23x, "학생")
_part = _bs.sheet_part["매크로작업"]
s23m = to_xlsm(s23x, os.path.join(TMP, "23학생코드만.xlsm"), VBA_BIN, _part,
               ("서식", "서식", "5, 0, 2, 0, 6, 0, 3, 0"))
bs = g.Book(s23m, "학생")
check(23, "xlsm의 양식 컨트롤 단추: 앵커·텍스트·연결 매크로 읽기",
      bs.form_controls("매크로작업") == [{"text": "서식", "anchor": "F3:G4",
                                      "macro": "서식"}],
      str(bs.form_controls("매크로작업")))
r, _, _ = grade(p23, a23, s23m)
check(23, "실행 결과가 없어도 VBA 코드에 지시 동작이 있으면 인정 -> 만점",
      r.earned == r.alloc, str(r.details) + str(r.notes))
check(23, "리포트에 발견한 매크로 이름·단추 앵커/텍스트/연결 매크로 표기",
      any("판매금액, 서식" in n for n in r.notes)
      and any("[F3:G4]" in n and "서식 매크로" in n for n in r.notes),
      str(r.notes))
check(23, "코드로 인정한 항목을 명시 + 실행까지 하라는 안내",
      any("매크로 코드에 해당 동작이 있어 인정" in n for n in r.notes),
      str(r.notes))

# 코드에도 없는 경우: 감점 유지 + xlsm 안내
NOOP_SRC = ('Attribute VB_Name = "Module1"\nSub 아무것도()\n'
            '    Range("Z90").Select\nEnd Sub\n'
            + "\n".join(f"' 여백 {i} " + "-" * 40 for i in range(80)))
s23n = to_xlsm(s23x, os.path.join(TMP, "23학생빈코드.xlsm"),
               build_vba_bin(NOOP_SRC), _part, None)
r, _, _ = grade(p23, a23, s23n)
check(23, "VBA 코드에도 지시 동작이 없으면 감점 유지", r.earned < r.alloc,
      str(r.details))
card = next((c for c in r.wrong if c["label"] == "매크로 서식 결과"), None)
check(23, "카드 설명에 '코드에 해당 문장이 없다'는 근거 표기",
      card is not None and "없었습니다" in str(card.get("note")),
      str(card.get("note")) if card else "카드 없음")

# 서식 값을 읽을 수 없는 경우: '판정 불가: 이유' 표기 (openpyxl 읽기 실패 재현)
check(23, "읽을 수 없는 서식 시그니처('?')는 값 대신 None -> 판정 불가로 이어짐",
      g._macro_fmt_desc("font", "?") is None
      and g._macro_fmt_desc("fill", None) == "채우기 없음")
_orig_desc = g._macro_fmt_desc
g._macro_fmt_desc = lambda kind, sig, raw=None: (
    None if kind == "font" else _orig_desc(kind, sig, raw))
try:
    r, _, _ = grade(p23, a23, s23n)
    card = next((c for c in r.wrong if c["label"] == "매크로 서식 결과"), None)
    props = [p2 for p2 in (card or {}).get("props", [])
             if p2["name"].startswith("글꼴")]
    check(23, "서식을 읽을 수 없으면 '판정 불가: 이유'로 명시",
          card is not None and props
          and "판정 불가:" in props[0]["expected"]
          and "판정 불가:" in props[0]["got"],
          str(props[:1]) if card else "카드 없음")
finally:
    g._macro_fmt_desc = _orig_desc

# --- 13-5 (선택) 배포 세트 정답 파일 자기 채점 100점 ---
SETS = os.path.normpath(os.path.join(BASE, "..", "모의고사"))
KEYS = os.path.normpath(os.path.join(BASE, "..", "기대값"))
found = 0
for n in ("1", "2"):
    prob = os.path.join(SETS, f"코코모의고사{n}회_문제.xlsx")
    ans = os.path.join(SETS, f"코코모의고사{n}회_정답.xlsx")
    keyp = os.path.join(KEYS, f"코코모의고사{n}회_기대값.json")
    if not (os.path.isfile(prob) and os.path.isfile(ans)):
        continue
    found += 1
    import json as _json
    kv = {}
    if os.path.isfile(keyp):
        with open(keyp, encoding="utf-8") as f:
            kv = _json.load(f)
    _res, score, _notes, _ = g.run_grading(prob, ans, ans, g.normalize_key(kv))
    lost = [(x.name, x.alloc, x.earned) for x in _res if x.earned < x.alloc]
    check(24, f"코코 모의고사 {n}회 정답 파일 자기 채점 100점 (정답 파일 잔여 서식 없음)",
          score == 100 and not lost, f"{score}점 {lost}")
if not found:
    print("  [24] 모의고사/ 세트 파일 없음 — 건너뜀")

# ---------------------------------------------------------------------------
print("25. v2.0.4 이의제기 2건 (2024 상시 2회, 코코 복원 문제지 기준)")
# ① 계산작업 3번 — 부등호 어순 교환은 동치, 둘째 조건의 참조(C열/D열)는 오답.
#    이 자료에서는 둘째 조건이 C열이어도 결과가 우연히 같다(본선 2명)는 점이 핵심.
# ② 매크로 서식 — [B5:F5], [B6:B10] 두 범위 중 머리글 행에만 적용한 풀이.

REC25 = [("존 스미스", 9.82, 9.79), ("마리아", 10.15, 10.25),
         ("아흐메드", 10.32, 10.45), ("리 웨이", 10.45, 10.42),
         ("엠마", 10.56, 10.61), ("송길동", 10.11, 10.14),
         ("이철수", 10.25, 10.33), ("마리안", 10.36, 10.52),
         ("안나", 10.40, 10.48)]
A25, B25 = 3, 3 + len(REC25) - 1
F_ANS25 = ('=IF(OR(C{r}<=SMALL($C$3:$C$11,2),D{r}<=SMALL($D$3:$D$11,2)),'
           '"본선","")')
F_SWAP25 = ('=IF(OR(SMALL($C$3:$C$11,2)>=C{r},SMALL($D$3:$D$11,2)>=D{r}),'
            '"본선","")')
F_BUG25 = ('=IF(OR(SMALL($C$3:$C$11,2)>=C{r},SMALL($D$3:$D$11,2)>=C{r}),'
           '"본선","")')


def wb_calc25(formula=None):
    wb = wb_new("계산작업")
    ws = wb.active
    ws["A1"] = "[표3] 국가대표 선수들의 달리기 경기 기록"
    for j, h in enumerate(["선수명", "국적", "1차 기록", "2차 기록", "결과"]):
        ws.cell(2, 1 + j).value = h
    for i, (nm, c1, c2) in enumerate(REC25):
        ws.cell(A25 + i, 1).value = nm
        ws.cell(A25 + i, 2).value = "국가"
        ws.cell(A25 + i, 3).value = c1
        ws.cell(A25 + i, 4).value = c2
        if formula:
            ws.cell(A25 + i, 5).value = formula.format(r=A25 + i)
    return wb


# 참조 횟수 비교 (단위 테스트)
check(25, "참조 횟수 비교: 둘째 조건이 C열이면 다른 참조로 인식",
      not g.same_ref_usage(F_ANS25.format(r=4), F_BUG25.format(r=4))
      and g.ref_usage_diff(F_ANS25.format(r=4),
                           F_BUG25.format(r=4)) == (["D4"], ["C4"]),
      str(g.ref_usage_diff(F_ANS25.format(r=4), F_BUG25.format(r=4))))
check(25, "참조 횟수 비교: 부등호 어순만 바꾼 식은 같은 참조",
      g.same_ref_usage(F_ANS25.format(r=4), F_SWAP25.format(r=4)))
check(25, "참조 횟수 비교: SUM 범위 표기와 나열 표기는 같은 참조",
      g.same_ref_usage("=SUM(C6:E6)", "=SUM(C6,D6,E6)")
      and g.same_ref_usage("=A3*2", "=2*A3"),
      str(g.formula_ref_counts("=SUM(C6:E6)")))
check(25, "참조 횟수 비교: 함수 이름(LOG10)은 참조로 세지 않음",
      g.formula_ref_counts("=LOG10(A1)+SUM(B2:B3)") ==
      {"A1": 1, "B2": 1, "B3": 1},
      str(g.formula_ref_counts("=LOG10(A1)+SUM(B2:B3)")))

p25 = save(wb_calc25(None), "25문제.xlsx")
a25 = save(wb_calc25(F_ANS25), "25정답.xlsx")
s25ok = save(wb_calc25(F_SWAP25), "25학생어순.xlsx")
s25bug = save(wb_calc25(F_BUG25), "25학생참조오류.xlsx")

if g.soffice_exe():
    r, _, _ = grade(p25, a25, s25ok)
    check(25, "계산값이 없어도 재계산해 값으로 비교 -> 어순만 바꾼 풀이는 만점",
          r.earned == r.alloc, str(r.details)[:200])
    check(25, "리포트에 '계산값이 없어 재계산해 비교했습니다' 명시",
          any("재계산해 비교했습니다" in n for n in r.notes), str(r.notes)[:160])

    r, _, _ = grade(p25, a25, s25bug)
    check(25, "재계산 결과가 정답과 같아도 둘째 조건의 참조가 다르면 오답 유지",
          r.earned < r.alloc, f"{r.earned}/{r.alloc} {r.details[:1]}")
    check(25, "오답 근거를 셀 단위로 표기 (정답은 D4, 내 답은 C4를 참조)",
          any("참조하는 셀이 정답과 달라" in n and "D4" in n and "C4" in n
              for n in r.notes),
          str([n for n in r.notes if "참조" in n])[:220])
    card25 = next((c for c in r.wrong if c["label"].startswith("계산 문제")), None)
    check(25, "카드에도 참조 차이 근거 + '우연히 결과가 같았을 뿐' 안내",
          card25 is not None and "참조" in (card25.get("note") or "")
          and "우연히" in (card25.get("note") or ""),
          str(card25.get("note"))[:200] if card25 else "카드 없음")
    check(25, "학생 셀 상태를 '수식만 있음 · 재계산값 …'으로 표기",
          card25 is not None
          and any(c.get("got_state") == "recalc" for c in card25["cells"]),
          str(card25["cells"][0]) if card25 else "카드 없음")

    _re = g.RECALC_ENABLED
    g.RECALC_ENABLED = False
    try:
        r, _, _ = grade(p25, a25, s25ok)
    finally:
        g.RECALC_ENABLED = _re
    check(25, "재계산을 끄면 조용히 수식 비교로 내려가고 그 사실을 표기",
          any("수식 문자열로 비교" in n for n in r.notes), str(r.notes)[:160])
else:
    print("  [25] LibreOffice 없음 — 재계산 경로 테스트 건너뜀")

# ② 매크로 -----------------------------------------------------------------
DEPT25 = [("소프트웨어 엔지니어", 84, 90, 82), ("시스템 분석가", 87, 95, 78),
          ("전략 개발", 83, 78, 90), ("마케팅", 79, 65, 89),
          ("생산 관리자", 77, 80, 95)]
BLUE25 = "0070C0"

MACRO25 = 'Attribute VB_Name = "Module1"\n' + "\n".join([
    "Sub 총점()",
    '    Range("F6").Select',
    '    ActiveCell.FormulaR1C1 = "=SUM(RC[-3]:RC[-1])"',
    '    Selection.AutoFill Destination:=Range("F6:F10")',
    "End Sub",
    "Sub 서식()",
    '    Range("B5:F5,B6:B10").Select',
    "    Selection.Font.Bold = True",
    "    With Selection.Font",
    "        .Color = -4165632",
    "    End With",
    "End Sub",
] + [f"' 여백 {i} " + "-" * 40 for i in range(80)])

MACRO25_HEAD = MACRO25.replace('Range("B5:F5,B6:B10").Select',
                               'Range("B5:F5").Select')

u25 = g.vba_macro_units(g.vba_module_sources(build_vba_bin(MACRO25))[0][1])
check(25, 'VBA 다중 영역 Range("B5:F5,B6:B10")를 영역별로 인식',
      g.vba_covers(u25, "B6", "font") == "서식"
      and g.vba_covers(u25, "F5", "font") == "서식"
      and g.vba_covers(u25, "C7", "font") is None,
      str([(x["name"], x["ranges"]) for x in u25]))


def wb_macro25(head=False, body=False):
    wb = wb_new("매크로작업")
    ws = wb.active
    ws["B3"] = "[표] 부서별 인사 평가"
    for j, h in enumerate(["부서", "기본자질", "업무지식", "의욕태도", "총점"]):
        ws.cell(5, 2 + j).value = h
    for i, (nm, a, b, c) in enumerate(DEPT25):
        r2 = 6 + i
        ws.cell(r2, 2).value = nm
        ws.cell(r2, 3).value = a
        ws.cell(r2, 4).value = b
        ws.cell(r2, 5).value = c
        ws.cell(r2, 6).value = f"=SUM(C{r2}:E{r2})"
    if head:
        for c in range(2, 7):
            ws.cell(5, c).font = Font(bold=True, color=BLUE25)
    if body:
        for i in range(len(DEPT25)):
            ws.cell(6 + i, 2).font = Font(bold=True, color=BLUE25)
    return wb


p25m = save(wb_macro25(), "25매크로문제.xlsx")
a25m = save(wb_macro25(head=True, body=True), "25매크로정답.xlsx")
s25m = save(wb_macro25(head=True), "25매크로머리글만.xlsx")

r, _, _ = grade(p25m, a25m, s25m)
card = next((c for c in r.wrong if c["label"] == "매크로 서식 결과"), None)
check(25, "두 범위 중 머리글 행에만 적용 -> 감점", r.earned < r.alloc,
      f"{r.earned}/{r.alloc}")
check(25, "맞은 범위(B5:F5)와 틀린 범위(B6:B10)를 함께 표기",
      card is not None and "B5:F5 영역은 정확히 적용" in (card.get("note") or "")
      and any("B6:B10" in p["name"] for p in card.get("props", [])),
      str(card.get("note"))[:200] if card else "카드 없음")
check(25, "글꼴 색을 사람이 읽는 이름으로 표기 (파랑(#0070C0) / 자동(검정))",
      card is not None
      and "파랑(#0070C0)" in str(card["cells"][0]["expected"])
      and "자동(검정)" in str(card["cells"][0]["got"]),
      str(card["cells"][0]) if card else "카드 없음")
check(25, "두 범위 모두 적용한 풀이는 만점",
      grade(p25m, a25m, save(wb_macro25(head=True, body=True),
                             "25매크로전부.xlsx"))[0].earned == r.alloc)

# VBA 근거 문구: 같은 서식 매크로가 있는데 범위만 다른 경우
_bs25 = g.Book(s25m, "학생")
_part25 = _bs25.sheet_part["매크로작업"]
s25mv = to_xlsm(s25m, os.path.join(TMP, "25매크로머리글만.xlsm"),
                build_vba_bin(MACRO25_HEAD), _part25, None)
r, _, _ = grade(p25m, a25m, s25mv)
card = next((c for c in r.wrong if c["label"] == "매크로 서식 결과"), None)
check(25, "매크로는 있으나 적용 범위가 다르면 '범위가 지시와 다릅니다'로 설명",
      card is not None
      and "적용 범위가 지시와 다릅니다" in (card.get("note") or "")
      and "B5:F5" in (card.get("note") or ""),
      str(card.get("note"))[:240] if card else "카드 없음")
a25mv = to_xlsm(a25m, os.path.join(TMP, "25매크로정답.xlsm"),
                build_vba_bin(MACRO25), g.Book(a25m, "정답")
                .sheet_part["매크로작업"], None)
r, _, _ = grade(p25m, a25mv, to_xlsm(
    save(wb_macro25(), "25매크로코드만.xlsx"),
    os.path.join(TMP, "25매크로코드만.xlsm"), build_vba_bin(MACRO25),
    g.Book(save(wb_macro25(), "25매크로코드만2.xlsx"), "x")
    .sheet_part["매크로작업"], None))
check(25, "다중 영역 매크로 코드가 있으면 실행 결과가 없어도 인정",
      any("매크로 코드에 해당 동작이 있어 인정" in n for n in r.notes),
      str(r.notes)[:200])

# 도형(육각형) 텍스트 대조 -------------------------------------------------
DRAW = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/'
        '2006/spreadsheetDrawing" xmlns:a="http://schemas.openxmlformats.org/'
        'drawingml/2006/main"><xdr:twoCellAnchor>'
        '<xdr:from><xdr:col>7</xdr:col><xdr:colOff>0</xdr:colOff>'
        '<xdr:row>7</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>'
        '<xdr:to><xdr:col>9</xdr:col><xdr:colOff>0</xdr:colOff>'
        '<xdr:row>9</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>'
        '<xdr:sp macro="[0]!서식" textlink="">'
        '<xdr:nvSpPr><xdr:cNvPr id="2" name="육각형 1"/>'
        '<xdr:cNvSpPr/></xdr:nvSpPr>'
        '<xdr:spPr><a:prstGeom prst="hexagon"><a:avLst/></a:prstGeom>'
        '</xdr:spPr><xdr:txBody><a:bodyPr/><a:p><a:r>'
        '<a:t>{text}</a:t></a:r></a:p></xdr:txBody></xdr:sp>'
        '<xdr:clientData/></xdr:twoCellAnchor></xdr:wsDr>')


def with_shape(src, out, sheet_part, text):
    """합성 통합 문서에 매크로 연결 도형(육각형) 1개를 넣는다."""
    import zipfile
    with zipfile.ZipFile(src) as zin:
        blobs = {n: zin.read(n) for n in zin.namelist()}
    blobs["xl/drawings/drawing1.xml"] = DRAW.format(text=text).encode("utf-8")
    ct = blobs["[Content_Types].xml"].decode("utf-8")
    if "drawing1.xml" not in ct:
        ct = ct.replace("</Types>",
                        '<Override PartName="/xl/drawings/drawing1.xml" '
                        'ContentType="application/vnd.openxmlformats-'
                        'officedocument.drawing+xml"/></Types>')
    blobs["[Content_Types].xml"] = ct.encode("utf-8")
    srels = re.sub(r"worksheets/(sheet\d+\.xml)$",
                   r"worksheets/_rels/\1.rels", sheet_part)
    cur = blobs.get(srels, (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/'
        '2006/relationships"></Relationships>').encode("utf-8")).decode("utf-8")
    blobs[srels] = cur.replace(
        "</Relationships>",
        '<Relationship Id="rIdDrw" Type="http://schemas.openxmlformats.org/'
        'officeDocument/2006/relationships/drawing" '
        'Target="/xl/drawings/drawing1.xml"/></Relationships>').encode("utf-8")
    sx = blobs[sheet_part].decode("utf-8")
    if "xmlns:r=" not in sx.split(">", 1)[0]:
        sx = sx.replace("<worksheet ", '<worksheet xmlns:r="http://'
                        'schemas.openxmlformats.org/officeDocument/2006/'
                        'relationships" ', 1)
    blobs[sheet_part] = sx.replace(
        "</worksheet>", '<drawing r:id="rIdDrw"/></worksheet>').encode("utf-8")
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for n, b in blobs.items():
            zout.writestr(n, b)
    return out


a25d = with_shape(a25m, os.path.join(TMP, "25정답도형.xlsx"), _part25, "서식적용")
_b = g.Book(a25d, "정답")
check(25, "도형(육각형)의 텍스트·연결 매크로·도형 종류를 읽는다",
      _b.drawing_shapes("매크로작업") ==
      [{"text": "서식적용", "anchor": "H8:I9", "macro": "서식",
        "geom": "hexagon", "name": "육각형 1", "hidden": False}],
      str(_b.drawing_shapes("매크로작업")))

s25ok_d = with_shape(save(wb_macro25(head=True, body=True), "25도형정답풀이.xlsx"),
                     os.path.join(TMP, "25학생도형OK.xlsx"), _part25, "서식적용")
s25bad_d = with_shape(save(wb_macro25(head=True, body=True), "25도형오타풀이.xlsx"),
                      os.path.join(TMP, "25학생도형오타.xlsx"), _part25, "서식전용")
r_ok, _, _ = grade(p25m, a25d, s25ok_d)
r_bad, _, _ = grade(p25m, a25d, s25bad_d)
check(25, "도형 텍스트가 지시와 같으면 감점 없음", r_ok.earned == r_ok.alloc,
      f"{r_ok.earned}/{r_ok.alloc} {r_ok.details}")
check(25, "도형 텍스트가 다르면 감점 + 정답/내 답 텍스트 표기",
      r_bad.earned < r_ok.earned
      and any("서식전용" in d and "서식적용" in d for d in r_bad.details),
      f"{r_bad.earned}/{r_bad.alloc} {r_bad.details}")
card = next((c for c in r_bad.wrong if c["label"] == "매크로 단추·도형"), None)
check(25, "단추·도형 카드에 정답 개체와 내 개체를 나란히 표기",
      card is not None and card["props"]
      and "서식적용" in card["props"][0]["expected"]
      and "서식전용" in card["props"][0]["got"],
      str(card["props"]) if card else "카드 없음")

# ---------------------------------------------------------------------------
print("26. v2.0.5 이의제기 3건 (코코 모의고사 2회, 자체 제작 문제지 기준)")
# ① 기본작업-2 열 너비 — Excel은 [열 너비] 대화상자 값에 좌우 여백 5픽셀을
#    더해 저장한다(맑은 고딕 11pt: 12 입력 -> 12.625 저장). 정답 파일이 raw 12면
#    정확히 12를 입력한 풀이가 오답 처리되던 문제.
# ② 계산작업 3번 — MID 결과는 '문자'라 숫자 머리글과 매칭되지 않아 *1이
#    필요하고, HLOOKUP의 행 번호는 2여야 한다. 진단이 실제 차이를 짚지 못하고
#    '$ 절대참조는 …' 일반 문구만 나오던 문제.
# ③ 계산작업 5번 — 함수 이름 오타(AVRAGE)는 #NAME?, 결과 뒤 &"명" 누락,
#    조건을 '">"&800000'으로 이어 붙인 것은 '">800000"'과 같은 뜻.

# --- 26-1 저장값 <-> 대화상자 문자 수 환산 (단위) ---
check(26, "Excel 저장값 계산: 12자 -> 맑은 고딕 12.625 / Calibri 12.7109375",
      g.col_width_stored(12, 8) == 12.625
      and g.col_width_stored(12, 7) == 12.7109375,
      f"{g.col_width_stored(12, 8)} / {g.col_width_stored(12, 7)}")
check(26, "저장값 12.625 · 12.7109375 · 12 모두 '12자'로 판정",
      all(g.col_width_is(v, 12) for v in (12.625, 12.7109375, 12.0)),
      str([(v, g.col_width_is(v, 12)) for v in (12.625, 12.7109375, 12.0)]))
check(26, "다른 너비는 그대로 오답 (8 입력=8.625 / raw 8 / 13.25)",
      not any(g.col_width_is(v, 12) for v in (8.625, 8.0, 13.25)),
      str([(v, g.col_width_is(v, 12)) for v in (8.625, 8.0, 13.25)]))
check(26, "리포트 표기 '12자(저장값 12.625)' / 패딩 없으면 '12자'",
      g.col_width_text(12.625, 12) == "12자(저장값 12.625)"
      and g.col_width_text(12.0, 12) == "12자"
      and g.col_width_text(None, 12) == "지정 안 함(기본 너비)",
      f"{g.col_width_text(12.625, 12)} | {g.col_width_text(12.0, 12)}")
check(26, "행 높이는 포인트라 패딩 없음 (28 저장값 = 28)",
      g.row_height_text(28) == "28" and g.row_height_text(None).startswith("지정"),
      g.row_height_text(28))


def wb_b2_26(width=None, height=None):
    wb = wb_new("기본작업-2")
    ws = wb.active
    ws["A1"] = "미래상사 사원 급여 지급 현황"
    for j, h in enumerate(("사원번호", "성명", "부서명")):
        ws.cell(3, 1 + j).value = h
    for r in range(4, 8):
        for c in range(1, 4):
            ws.cell(r, c).value = f"v{r}{c}"
    if width is not None:
        ws.column_dimensions["A"].width = width
    if height is not None:
        ws.row_dimensions[1].height = height
    return wb


KEY26 = {"format_checks": {"기본작업-2": [
    {"no": 1, "range": "1행", "check": {"row_height": 28},
     "지시": "1행 높이를 28로 지정"},
    {"no": 2, "range": "A열", "check": {"col_width": 12},
     "지시": "A열 너비를 12로 지정"}]}}

p26 = save(wb_b2_26(), "26문제.xlsx")
a26 = save(wb_b2_26(width=12, height=28), "26정답raw.xlsx")
a26x = save(wb_b2_26(width=12.625, height=28), "26정답excel.xlsx")

for tag, stu_w in (("맑은 고딕 12.625", 12.625),
                   ("Calibri 12.7109375", 12.7109375),
                   ("그대로 저장된 12", 12.0)):
    s = save(wb_b2_26(width=stu_w, height=28), f"26학생{stu_w}.xlsx")
    for atag, ap in (("정답 raw 12", a26), ("정답 Excel 12.625", a26x)):
        r, _, _ = grade(p26, ap, s, key=KEY26)
        check(26, f"열 너비 '12' 입력({tag}) × {atag} -> 만점",
              r.earned == r.alloc, f"{r.earned}/{r.alloc} {r.details}")

s26bad = save(wb_b2_26(width=8.625, height=28), "26학생너비오답.xlsx")
r, _, _ = grade(p26, a26, s26bad, key=KEY26)
card26 = next((c for c in r.wrong if c["label"] == "서식 - 열 너비"), None)
check(26, "정말 다른 너비(8 입력)는 감점 + 'A열' 표기",
      r.earned < r.alloc and any("A열" in d for d in r.details),
      f"{r.earned}/{r.alloc} {r.details}")
check(26, "카드에 '정답 12자 / 내 답 8자(저장값 8.625)'로 사람이 읽게 표기",
      card26 is not None
      and card26["props"][0]["expected"] == "12자"
      and card26["props"][0]["got"] == "8자(저장값 8.625)",
      str(card26["props"]) if card26 else "카드 없음")
check(26, "카드 노트에 '저장' 방식(여백 5픽셀) 설명",
      card26 is not None and "5픽셀" in (card26.get("note") or ""),
      str(card26.get("note"))[:80] if card26 else "카드 없음")

s26h = save(wb_b2_26(width=12.625, height=28.5), "26학생행높이반올림.xlsx")
r, _, _ = grade(p26, a26, s26h, key=KEY26)
check(26, "행 높이 1픽셀(0.75pt) 반올림(28.5)은 통과", r.earned == r.alloc,
      f"{r.earned}/{r.alloc} {r.details}")
s26hb = save(wb_b2_26(width=12.625, height=20), "26학생행높이오답.xlsx")
r, _, _ = grade(p26, a26, s26hb, key=KEY26)
check(26, "행 높이가 실제로 다르면(20) 감점 + 정답/내 답 표기",
      r.earned < r.alloc
      and any("행 높이" in d and "28" in d and "20" in d for d in r.details),
      f"{r.earned}/{r.alloc} {r.details}")

# --- 26-2 계산작업 진단: 문자 -> 숫자(*1), HLOOKUP 행 번호 ---
F_ANS26 = '=HLOOKUP(MID(A{r},2,1)*1,$B$9:$D$10,2,TRUE)'
F_STU26 = '=HLOOKUP(MID(A{r},2,1),$B$9:$D$10,1,TRUE)'
dn26 = g.diagnose_formula_diff(F_STU26.format(r=3), F_ANS26.format(r=3))
check(26, "진단: MID 결과는 문자라 *1로 숫자로 바꿔야 한다고 설명",
      any("문자" in n and "*1" in n for n in dn26), str(dn26))
check(26, "진단: HLOOKUP 3번째 인수를 '가져올 행 번호'로 지적(2 vs 1)",
      any("가져올 행 번호" in n and "2" in n for n in dn26), str(dn26))
check(26, "진단에 '$ 절대참조' 일반 문구만 남지 않는다",
      dn26 and not all("절대참조" in n or "$" in n for n in dn26), str(dn26))
pt26 = g.formula_point(F_ANS26.format(r=3), F_STU26.format(r=3))
check(26, "포인트가 '$ 절대참조는 …'이 아니라 문자↔숫자 설명",
      pt26 is not None and "문자" in pt26 and "VALUE" in pt26
      and pt26 != g.ABS_POINT, str(pt26)[:90])
check(26, "*1 을 제대로 쓴 풀이에는 문자↔숫자 진단이 안 붙는다",
      not any("*1" in n and "문자" in n for n in
              g.diagnose_formula_diff(
                  '=HLOOKUP(MID(A3,2,1)*1,$B$9:$D$10,1,TRUE)',
                  F_ANS26.format(r=3))),
      str(g.diagnose_formula_diff(
          '=HLOOKUP(MID(A3,2,1)*1,$B$9:$D$10,1,TRUE)', F_ANS26.format(r=3))))

# --- 26-3 계산작업 진단: 함수 이름 오타 · 조건 문자열 · &"명" ---
F_ANS26B = '=COUNTIFS(D31:D38,">800000",C31:C38,">="&AVERAGE(C31:C38))&"명"'
F_STU26B = '=COUNTIFS(D31:D38,">"&800000,C31:C38,">="&AVRAGE(C31:C38))'
dn26b = g.diagnose_formula_diff(F_STU26B, F_ANS26B)
check(26, "진단: 함수 이름 오타(AVRAGE → AVERAGE)와 #NAME? 안내",
      any("오타" in n and "AVRAGE" in n and "AVERAGE" in n and "#NAME?" in n
          for n in dn26b), str(dn26b))
check(26, '진단: \'">"&800000\' 은 \'">800000"\' 과 같은 뜻이라고 인정',
      any("같은 뜻" in n and "맞습니다" in n for n in dn26b), str(dn26b))
check(26, '진단: 결과 뒤 &"명" 누락을 지적',
      any("명" in n and "&" in n and "이어 붙이기가 빠졌습니다" in n for n in dn26b),
      str(dn26b))
pt26b = g.formula_point(F_ANS26B, F_STU26B)
check(26, "포인트: 상수는 따옴표 안에, 함수 결과는 & 로 이어 붙인다",
      pt26b is not None and "따옴표" in pt26b and "AVERAGE" in pt26b,
      str(pt26b)[:90])
check(26, '&"명" 까지 쓴 풀이에는 누락 진단이 안 붙는다',
      not any("이어 붙이기가 빠졌습니다" in n for n in g.diagnose_formula_diff(
          '=COUNTIFS(D31:D38,">"&800000,C31:C38,">="&AVERAGE(C31:C38))&"명"',
          F_ANS26B)),
      str(g.diagnose_formula_diff(
          '=COUNTIFS(D31:D38,">"&800000,C31:C38,">="&AVERAGE(C31:C38))&"명"',
          F_ANS26B)))

# --- 26-4 오류 값은 0으로 뭉개지 않고 그대로 보여 준다 ---
check(26, "Excel 오류 값 인식과 원인 한 줄",
      g.excel_error("#NAME?") == "#NAME?" and g.excel_error("3명") is None
      and g.excel_error(0) is None
      and "함수 이름" in (g.excel_error_note("#NAME?") or ""),
      str(g.excel_error_note("#NAME?"))[:70])


def wb_calc26(mode):
    """mode: 'problem' | 'answer' | 'error'(학생 셀에 #NAME? 결과)"""
    wb = wb_new("계산작업")
    ws = wb.active
    ws["A1"] = "[표5] 코코몰 주문 현황"
    for j, h in enumerate(("고객명", "구매횟수", "주문금액")):
        ws.cell(2, 1 + j).value = h
    for i in range(3, 11):
        ws.cell(i, 1).value = f"고객{i}"
        ws.cell(i, 2).value = 8 + i
        ws.cell(i, 3).value = 700000 + i * 50000
    ws["A12"] = "우수 고객 수"
    if mode == "answer":
        ws["D12"] = F_ANS26B.replace("D31:D38", "C3:C10").replace(
            "C31:C38", "B3:B10")
    elif mode == "error":
        ws["D12"] = "#NAME?"          # 오타 수식이 남긴 오류 값
    return wb


p26c = save(wb_calc26("problem"), "26계산문제.xlsx")
a26c = save(wb_calc26("answer"), "26계산정답.xlsx")
s26c = save(wb_calc26("error"), "26계산오류학생.xlsx")
r, _, _ = grade(p26c, a26c, s26c)
card26c = next((c for c in r.wrong if c["label"].startswith("계산 문제")), None)
check(26, "오류 값은 숫자 0이 아니라 '#NAME? (오류 값)'으로 표기",
      card26c is not None
      and card26c["cells"][0]["got"] == "#NAME? (오류 값)"
      and card26c["cells"][0]["got_state"] == "error",
      str(card26c["cells"][0]) if card26c else "카드 없음")
check(26, "카드 노트에 오류 원인 한 줄",
      card26c is not None and "#NAME?" in (card26c.get("note") or "")
      and "함수 이름" in (card26c.get("note") or ""),
      str(card26c.get("note"))[-90:] if card26c else "카드 없음")

# ---------------------------------------------------------------------------
print()
print("27. v2.0.6 이의제기 3건 (컴활 2급 상시 = 사용자 '상시기출2회' 리포트)")
# ① 기본작업-1 띄어쓰기 — 복원 문제지가 자료 표를 전부 가운데 정렬로 그려
#    공백을 사람 눈으로 분간할 수 없던 문제(문제지 쪽 결함, 점수는 정정 데이터로
#    되돌림). 채점기 쪽 재현 대상은 아니므로 여기서는 다루지 않는다.
# ② 기본작업-2 — (a) '아래쪽 테두리' 지시가 늘 아래 셀(A2 위쪽)로 표기돼
#    지시 범위 밖 주소를 인용하던 문제 (b) 정답 파일에만 있는 제작 메모
#    ('區分+E4:E13')를 셀 값 지시로 보고 감점하던 문제.
# ③ 매크로작업 — 4096바이트 섹터(CFB v4) vbaProject.bin에서 스트림을 하나도
#    못 읽어 VBA 기반 판정이 조용히 죽어 있던 문제.

# --- 27-1 CFB v4(4096바이트 섹터) vbaProject.bin 읽기 ---
SRC27 = 'Attribute VB_Name = "Module1"\n' + "\n".join([
    "Sub 통화()",
    '    Range("E4:E11").Select',
    '    Selection.NumberFormatLocal = "\\#,##0_);[빨강](\\#,##0)"',
    "End Sub",
]) + "\n" + "\n".join(f"' 여백 {i} " + "-" * 40 for i in range(80))

for ssz in (512, 4096):
    binv = build_vba_bin(SRC27, ssz=ssz)
    mods = g.vba_module_sources(binv)
    check(27, f"{ssz}바이트 섹터 vbaProject.bin에서 모듈 소스를 읽는다",
          len(mods) == 1 and "Sub 통화" in mods[0][1],
          f"모듈 {len(mods)}개")
    units = g.vba_macro_units(mods[0][1]) if mods else []
    check(27, f"{ssz}바이트 섹터: 매크로 이름·범위·동작 인식",
          len(units) == 1 and units[0]["name"] == "통화"
          and "E4:E11" in units[0]["ranges"]
          and "number_format" in units[0]["kinds"],
          str(units))
    check(27, f"{ssz}바이트 섹터: 읽기 진단이 '정상'(None)",
          g.vba_read_diag(binv) is None, str(g.vba_read_diag(binv)))

check(27, "리포트 ID = 응시 시각(기록 id와 같음), 풀이 이름에 없으면 None",
      g.attempt_report_id("풀이_컴활 2급 상시_20260915_1226.xlsm")
      == "20260915122600"
      and g.attempt_report_id("풀이_세트_20260915_1226 (2).xlsm")
      == "20260915122600"
      and g.attempt_report_id("정답.xlsx") is None,
      str(g.attempt_report_id("풀이_세트_20260915_1226 (2).xlsm")))
check(27, "CFB가 아니면 사유를 'OLE(CFB) 형식이 아닙니다'로 알린다",
      "OLE(CFB) 형식이" in (g.vba_read_diag(b"PK\x03\x04" + b"\x00" * 600) or ""),
      str(g.vba_read_diag(b"PK\x03\x04" + b"\x00" * 600)))
check(27, "구조는 CFB인데 깨졌으면 '구조를 해석하지 못했습니다'",
      "구조를 해석하지" in (g.vba_read_diag(CFB_MAGIC + b"\x00" * 600) or ""),
      str(g.vba_read_diag(CFB_MAGIC + b"\x00" * 600)))


# --- 27-1b 단추 캡션: Excel 이 글자 뒤에 붙이는 <br/> 때문에 못 읽던 버그 ---
VML_BTN = ('<v:shape id="_x0000_s8193" type="#_x0000_t201" o:button="t">'
           '<v:textbox style=\'mso-direction-alt:auto\' o:singleclick="f">'
           '<div style=\'text-align:center\'><font face="\ub9d1\uc740 \uace0\ub515" '
           'size="220" color="#000000">{cap}</font></div></v:textbox>'
           '<x:ClientData ObjectType="Button"><x:Anchor>7, 0, 3, 0, 9, 0, 5, 0'
           '</x:Anchor><x:FmlaMacro>[0]!\ud3c9\uade0</x:FmlaMacro>'
           '</x:ClientData></v:shape>')
check(27, "단추 캡션 뒤 <br/> 가 있어도 텍스트를 읽는다 (Excel 이 실제로 쓰는 모양)",
      g._vml_textbox_text(VML_BTN.format(cap="평균<br\n   />\n      ")) == "평균",
      repr(g._vml_textbox_text(VML_BTN.format(cap="평균<br\n   />\n      "))))
check(27, "<br/> 없는 평범한 캡션도 그대로 읽는다",
      g._vml_textbox_text(VML_BTN.format(cap="서식")) == "서식",
      repr(g._vml_textbox_text(VML_BTN.format(cap="서식"))))
check(27, "여러 줄 캡션은 공백 한 칸으로 이어 붙인다",
      g._vml_textbox_text(VML_BTN.format(cap="평균<br/>계산")) == "평균 계산",
      repr(g._vml_textbox_text(VML_BTN.format(cap="평균<br/>계산"))))
check(27, "캡션이 없으면 빈 문자열",
      g._vml_textbox_text('<v:shape o:button="t"><v:textbox><div></div>'
                          '</v:textbox></v:shape>') == "",
      repr(g._vml_textbox_text('<v:shape o:button="t"><v:textbox><div></div>'
                               '</v:textbox></v:shape>')))
check(27, "&amp; 같은 이스케이프는 원래 글자로",
      g._vml_textbox_text(VML_BTN.format(cap="합계&amp;평균<br/>")) == "합계&평균",
      repr(g._vml_textbox_text(VML_BTN.format(cap="합계&amp;평균<br/>"))))


# --- 27-2 매크로 범위가 지시와 다르면 내 VBA 코드의 실제 범위를 인용 ---
def wb_macro27(applied_cols):
    """E·F열 통화 서식 매크로. applied_cols: 학생이 실제로 적용한 열."""
    wb = wb_new("매크로작업")
    ws = wb.active
    for j, h in enumerate(("상품명", "단가", "수량", "판매금액")):
        ws.cell(3, 3 + j).value = h
    for r in range(4, 12):
        ws.cell(r, 3).value = f"상품{r}"
        ws.cell(r, 4).value = 1000 * r
        ws.cell(r, 5).value = 1000 * r
        ws.cell(r, 6).value = 2000 * r
    for c in applied_cols:
        for r in range(4, 12):
            ws.cell(r, c).number_format = '"\u20a9"#,##0_);[Red]\\("\u20a9"#,##0\\)'
    return wb


p27 = save(wb_macro27(()), "27문제.xlsx")
a27 = save(wb_macro27((5, 6)), "27정답.xlsx")
s27 = save(wb_macro27((5,)), "27학생.xlsx")
P27 = os.path.join(TMP, "27문제.xlsm")
A27 = os.path.join(TMP, "27정답.xlsm")
S27 = os.path.join(TMP, "27학생.xlsm")
to_xlsm(p27, P27, build_vba_bin(SRC27, ssz=4096))
to_xlsm(a27, A27, build_vba_bin(
    SRC27.replace('Range("E4:E11")', 'Range("E4:F11")'), ssz=4096),
    button=("통화", "통화", "1, 0, 12, 0, 3, 0, 14, 0"))
to_xlsm(s27, S27, build_vba_bin(SRC27, ssz=4096),
        button=("통화", "통화", "1, 0, 12, 0, 3, 0, 14, 0"))
r27, _, _ = grade(P27, A27, S27)
note27 = " ".join(str(c.get("note") or "") for c in r27.wrong)
check(27, "적용 안 한 F열만 감점 (E열은 통과)",
      r27.earned < r27.alloc and "F4:F11" in str(r27.details),
      str(r27.details))
check(27, "리포트가 내 VBA 코드의 실제 범위 [E4:E11]를 인용",
      "E4:E11" in note27 and "적용 범위가 지시와 다릅니다" in note27,
      note27[-140:])
check(27, "참고 노트에 내 매크로 이름('통화')이 남는다",
      any("통화" in n for n in r27.notes), str(r27.notes)[:160])

# VBA가 깨져 못 읽히면 조용히 넘어가지 않고 사유를 적는다
S27B = os.path.join(TMP, "27학생_깨진VBA.xlsm")
to_xlsm(s27, S27B, CFB_MAGIC + b"\x00" * 4096,
        button=("통화", "통화", "1, 0, 12, 0, 3, 0, 14, 0"))
r27b, _, _ = grade(P27, A27, S27B)
all27b = " ".join(r27b.notes) + " " + " ".join(
    str(c.get("note") or "") for c in r27b.wrong)
check(27, "VBA를 못 읽으면 'VBA를 읽지 못했습니다' + 사유를 리포트에 남긴다",
      "VBA를 읽지 못했습니다" in all27b and "구조를 해석하지" in all27b,
      all27b[:160])


# --- 27-3 아래쪽 테두리 지시는 지시 범위 안 주소로 표기 ---
def wb_b2_27(bottom=False, memo=False, stray=False):
    wb = wb_new("기본작업-2")
    ws = wb.active
    ws["A1"] = "연계전공 가능 전공"
    for j, h in enumerate(("과목코드", "과목명", "구분", "모집인원",
                           "수업요일", "학점", "담당교수")):
        ws.cell(3, 1 + j).value = h
    for r in range(4, 14):
        for c in range(1, 8):
            ws.cell(r, c).value = f"v{r}{c}"
    if bottom:
        for c in range(1, 8):
            ws.cell(1, c).border = Border(bottom=Side(style="thick"))
    if memo:
        ws["E2"] = "區分+E4:E13"
    if stray:
        ws["E2"] = "구분"
    return wb


p27b = save(wb_b2_27(), "27b문제.xlsx")
a27b = save(wb_b2_27(bottom=True), "27b정답.xlsx")
s27b = save(wb_b2_27(), "27b학생.xlsx")
r27b2, _, _ = grade(p27b, a27b, s27b)
det27 = str(r27b2.details) + str([c.get("props") for c in r27b2.wrong])
check(27, "'A1:G1 아래쪽 테두리' 감점을 지시 범위 안 주소(A1)로 표기",
      "A1:G1" in det27 and "A2:G2" not in det27, det27[:200])

# 제작 메모(범위 참조 포함)는 채점 제외, 범위 참조 없는 정답 전용 값은 감점 유지
a27c = save(wb_b2_27(bottom=True, memo=True), "27c정답.xlsx")
r27c, _, _ = grade(p27b, a27c, s27b)
notes27c = " ".join(r27c.notes)
check(27, "정답 파일에만 있는 제작 메모 'X+E4:E13'은 채점에서 제외",
      "E2" not in str(r27c.details) and "제작 메모" in notes27c,
      notes27c[:150])
check(27, "제작 메모를 뺀 뒤에도 테두리 감점은 그대로",
      "A1:G1" in str(r27c.details), str(r27c.details)[:150])

a27d = save(wb_b2_27(bottom=True, stray=True), "27d정답.xlsx")
r27d, _, _ = grade(p27b, a27d, s27b)
check(27, "범위 참조가 없는 정답 전용 값('구분')은 그대로 감점",
      "E2" in str(r27d.details), str(r27d.details)[:150])




# ---------------------------------------------------------------------------
print()
print("28. v2.0.8 모르면 감점 금지 — UNKNOWN 보류 · 확인 필요 표기 · "
      "실제 Excel 마크업")

# --- 28-1 UNKNOWN 센티널 자체 ---
U = g.Unknown("테스트 사유")
check(28, "Unknown 은 falsy·빈 컨테이너라 기존 소비자 코드가 그대로 돈다",
      (not U) and len(U) == 0 and list(U) == [] and U.get("k") is None
      and ("k" not in U) and list(U.items()) == [],
      repr(U))
check(28, "Unknown 끼리도 '같다'로 판정되지 않는다 (모르는 값끼리 우연 통과 금지)",
      g.Unknown("a") != g.Unknown("a") and U == U and g.is_unknown(U)
      and not g.is_unknown([]) and not g.is_unknown(None),
      "")
check(28, "unknown_why 가 사유를 그대로 돌려준다",
      g.unknown_why([], None, U) == "테스트 사유"
      and g.unknown_why([], None) == "원인을 특정하지 못했습니다",
      g.unknown_why(U))

# --- 28-2 추출 함수: '없다'와 '못 읽었다'를 가른다 ---
check(28, "cfb_streams: CFB 가 아니면 Unknown('형식이 아닙니다')",
      g.is_unknown(g.cfb_streams(b"PK\x03\x04" + b"\x00" * 600)),
      g.cfb_streams(b"PK\x03\x04" + b"\x00" * 600).why)
broken = CFB_MAGIC + b"\x00" * 4096
check(28, "cfb_streams: 헤더는 CFB인데 구조가 깨졌으면 Unknown (예전엔 {})",
      g.is_unknown(g.cfb_streams(broken)), g.cfb_streams(broken).why)
check(28, "vba_module_sources: 못 읽으면 Unknown (예전엔 [] = '매크로 없음')",
      g.is_unknown(g.vba_module_sources(broken)),
      g.vba_module_sources(broken).why)
check(28, "정상 vbaProject.bin 은 Unknown 이 아니다 (보류 남발 금지)",
      not g.is_unknown(g.vba_module_sources(VBA_BIN))
      and len(g.vba_module_sources(VBA_BIN)) == 1, "")
check(28, "fmt_signature: 셀을 못 읽으면 Unknown (예전 '?' 는 '?'끼리 같다고 "
          "판정돼 근거 없이 통과했다)",
      g.is_unknown(g.fmt_signature(None, "font"))
      and g.fmt_signature(None, "font") != g.fmt_signature(None, "font"),
      "")


class _BoomWS:
    """조건부 서식/병합을 읽는 순간 터지는 가짜 시트."""

    @property
    def conditional_formatting(self):
        raise RuntimeError("깨진 파트")

    @property
    def merged_cells(self):
        raise RuntimeError("깨진 파트")


check(28, "cf_rules_detail: 규칙을 읽다 실패하면 부분 목록이 아니라 Unknown",
      g.is_unknown(g.cf_rules_detail(_BoomWS())),
      g.cf_rules_detail(_BoomWS()).why)
_bm, _ba = g.merge_maps(_BoomWS())
check(28, "merge_maps: 병합을 못 읽으면 Unknown (서식 비교 기준이 틀어진다)",
      g.is_unknown(_bm) and g.is_unknown(_ba), g.unknown_why(_bm))


class _BoomBook:
    label = "가짜"

    class raw:
        @staticmethod
        def __getitem__(k):
            raise KeyError(k)

        class defined_names:
            @staticmethod
            def items():
                raise RuntimeError("깨진 파트")


check(28, "defined_names_for_sheet: 이름 정의를 못 읽으면 Unknown",
      g.is_unknown(g.defined_names_for_sheet(_BoomBook(), "기본작업-2")),
      g.defined_names_for_sheet(_BoomBook(), "기본작업-2").why)


# --- 28-3 보류가 감점으로 흐르지 않는가 (통합) ---
def wb_mac28(applied=False):
    wb = wb_new("매크로작업")
    ws = wb.active
    for r in range(3, 7):
        for c in range(1, 5):
            ws.cell(r, c).value = r * c
    if applied:
        for c in range(1, 5):
            ws.cell(3, c).fill = PatternFill("solid", fgColor="FFFF00")
    return wb


P28 = save(wb_mac28(), "28문제.xlsx")
A28 = save(wb_mac28(True), "28정답.xlsx")
S28 = save(wb_mac28(), "28학생.xlsx")          # 서식 미적용
S28X = to_xlsm(S28, os.path.join(TMP, "28학생_깨진VBA.xlsm"), broken)
r28, sc28, nt28 = grade(P28, A28, S28X)
check(28, "매크로 코드를 못 읽으면 결과 불일치를 감점하지 않고 보류한다",
      len(r28.holds) == 1 and r28.held == 8.0 and not r28.wrong,
      f"holds={len(r28.holds)} held={r28.held} 카드={len(r28.wrong)}")
check(28, "보류 사유에 '무엇을 왜' 가 적힌다",
      "vbaProject.bin" in r28.holds[0]["why"]
      and "A3(채우기)" in r28.holds[0]["why"], r28.holds[0]["why"][:90])
check(28, "보류분은 만점에서 빠져 점수는 남은 배점 기준으로 환산된다",
      sc28 == 100 and r28.earned == 2, f"{sc28} / earned={r28.earned}")

# 실제 Excel 은 4096바이트 미만 모듈 스트림을 미니 스트림에 담는다 —
# 매크로 모듈은 거의 다 그보다 작으니 실물은 대부분 이 경로를 탄다.
MINI_BIN = build_vba_bin_mini('Attribute VB_Name = "Module1"\n'
                              "Sub 아무것도()\nEnd Sub\n")
check(28, "실제 Excel 배치(미니 스트림, 64바이트 미니 섹터)의 "
          "vbaProject.bin 을 읽는다",
      g.vba_read_diag(MINI_BIN) is None
      and [u["name"] for u in g.vba_macro_units(
          g.vba_module_sources(MINI_BIN)[0][1])] == ["아무것도"],
      str(g.vba_read_diag(MINI_BIN)))

# 같은 학생 파일인데 VBA 가 멀쩡하면 (= 코드에도 그 동작이 없으면) 감점 유지
S28OK = to_xlsm(S28, os.path.join(TMP, "28학생_정상VBA.xlsm"), MINI_BIN)
r28b, sc28b, _ = grade(P28, A28, S28OK)
check(28, "미탐 방지: VBA 를 읽을 수 있으면 서식 미적용은 그대로 감점",
      not r28b.holds and r28b.earned < r28b.alloc and r28b.wrong,
      f"holds={len(r28b.holds)} earned={r28b.earned}/{r28b.alloc}")


def inject_rels(src, out, vml_body=None, target="/xl/drawings/vmlDrawing1.vml",
                sheet_part="xl/worksheets/sheet1.xml"):
    """시트 rels 에 vmlDrawing 을 선언한다 (vml_body=None 이면 파트를 안 넣는다)."""
    import zipfile
    with zipfile.ZipFile(src) as zin:
        blobs = {n: zin.read(n) for n in zin.namelist()}
    srels = re.sub(r"worksheets/(sheet\d+\.xml)$",
                   r"worksheets/_rels/\1.rels", sheet_part)
    blobs[srels] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/'
        '2006/relationships"><Relationship Id="rIdVml" Type="http://schemas.'
        'openxmlformats.org/officeDocument/2006/relationships/vmlDrawing" '
        f'Target="{target}"/></Relationships>').encode("utf-8")
    sx = blobs[sheet_part].decode("utf-8")
    if "xmlns:r=" not in sx.split(">", 1)[0]:
        sx = sx.replace("<worksheet ", '<worksheet xmlns:r="http://schemas.'
                        'openxmlformats.org/officeDocument/2006/'
                        'relationships" ', 1)
    blobs[sheet_part] = sx.replace(
        "</worksheet>", '<legacyDrawing r:id="rIdVml"/></worksheet>'
    ).encode("utf-8")
    if vml_body is not None:
        blobs["xl/drawings/vmlDrawing1.vml"] = vml_body.encode("utf-8")
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for n, b in blobs.items():
            zout.writestr(n, b)
    return out


A28V = inject_rels(save(wb_mac28(True), "28정답b.xlsx"),
                   os.path.join(TMP, "28정답_vml.xlsx"), "<xml/>")
S28V = inject_rels(save(wb_mac28(True), "28학생b.xlsx"),
                   os.path.join(TMP, "28학생_vml없음.xlsx"), None)
r28c, _, _ = grade(P28, A28V, S28V)
check(28, "rels 가 가리키는 VML 파트를 못 읽으면 '단추 없음 -2점'이 아니라 보류",
      any("단추" in h["label"] for h in r28c.holds) and r28c.held == 2.0
      and not any("단추" in c["label"] for c in r28c.wrong),
      f"holds={[h['label'] for h in r28c.holds]}")


# 조건부 서식 규칙을 못 읽으면 기본작업-3 전체 보류
def wb_cf28(rule=False):
    wb = wb_new("기본작업-3")
    ws = wb.active
    for r in range(3, 9):
        for c in range(1, 4):
            ws.cell(r, c).value = r * c
    if rule:
        from openpyxl.formatting.rule import FormulaRule
        ws.conditional_formatting.add(
            "A3:C8", FormulaRule(formula=["$C3>10"],
                                 font=Font(bold=True)))
    return wb


P28C = save(wb_cf28(), "28c문제.xlsx")
A28C = save(wb_cf28(True), "28c정답.xlsx")
S28C = save(wb_cf28(), "28c학생.xlsx")
_orig_cf = g.cf_rules_detail
try:
    def _cf_all_blind(ws, wb=None):
        return g.Unknown("조건부 서식 규칙을 읽지 못했습니다(테스트)")
    g.cf_rules_detail = _cf_all_blind
    r28d, sc28d, _ = grade(P28C, A28C, S28C)
    check(28, "조건부 서식 규칙을 못 읽으면 기본작업-3 을 감점하지 않고 보류",
          r28d.held == r28d.alloc and not r28d.wrong
          and "조건부 서식" in r28d.holds[0]["why"],
          f"held={r28d.held}/{r28d.alloc} why={r28d.holds[0]['why'][:50]}")
finally:
    g.cf_rules_detail = _orig_cf

# 미탐 방지: 규칙을 읽을 수 있으면 규칙 없는 학생은 그대로 감점
r28e, _, _ = grade(P28C, A28C, S28C)
check(28, "미탐 방지: 규칙을 읽을 수 있으면 조건부 서식 누락은 그대로 감점",
      not r28e.holds and r28e.earned < r28e.alloc,
      f"holds={len(r28e.holds)} earned={r28e.earned}/{r28e.alloc}")

# 채점기가 죽으면 그 시트 전체를 보류 (값 비교 결과만 남겨 감점하지 않는다)
_orig_b2 = g.GRADERS["기본작업-2"]
try:
    def _boom(res, ctx):
        raise RuntimeError("테스트용 강제 예외")
    g.GRADERS["기본작업-2"] = _boom
    P28F = save(wb_b2_27(), "28f문제.xlsx")
    A28F = save(wb_b2_27(bottom=True), "28f정답.xlsx")
    S28F = save(wb_b2_27(), "28f학생.xlsx")
    r28f, _, _ = grade(P28F, A28F, S28F)
    check(28, "채점기가 예외로 죽으면 그 시트 배점을 통째로 보류한다",
          r28f.held == r28f.alloc and r28f.earned == 0
          and "채점기가 이 시트를" in r28f.holds[0]["why"],
          f"held={r28f.held}/{r28f.alloc}")
finally:
    g.GRADERS["기본작업-2"] = _orig_b2


# --- 28-4 확인 필요 표기 (콘솔 / JSON / HTML) ---
import contextlib as _ctx
import io as _io
import json
_buf = _io.StringIO()
with _ctx.redirect_stdout(_buf):
    _res28, _sc28, _nt28, _ = g.run_grading(P28, A28, S28X, g.normalize_key({}))
    g.print_console(_res28, _sc28, _nt28, (P28, A28, S28X))
_con = _buf.getvalue()
check(28, "콘솔 총점 옆에 '88점 · 확인 필요 N건' 꼴로 적는다",
      f"{_sc28}점 · 확인 필요 1건" in _con, _con.split("\n")[13][:60])
check(28, "콘솔에 [확인 필요] 섹션과 계산 방식이 적힌다",
      "[확인 필요]" in _con and "만점에서 뺐습니다" in _con
      and "채점된 배점" in _con, "")

_j28 = os.path.join(TMP, "28결과.json")
g.write_json(_j28, _res28, _sc28, _nt28, (P28, A28, S28X))
with open(_j28, encoding="utf-8") as f:
    _d28 = json.load(f)
check(28, "JSON 기존 키는 그대로 (시험장.py 호환)",
      all(k in _d28 for k in ("total", "pass_line", "passed", "mode",
                              "graded_sheets", "max_total", "generated",
                              "files", "sheets", "notes", "wrong_items"))
      and all(k in _d28["sheets"][0] for k in ("name", "alloc", "earned",
                                               "missing", "details", "notes")),
      "")
check(28, "JSON 에 확인 필요 목록·배점·계산 방식이 추가된다",
      _d28["review_count"] == 1 and _d28["review_points"] == 8.0
      and _d28["review_items"][0]["sheet"] == "매크로작업"
      and _d28["scoring"]["gradable_total"] == 2.0
      and "만점에서 빼고" in _d28["scoring"]["method"]
      and _d28["sheets"][0]["held"] == 8.0,
      f"review_count={_d28['review_count']}")

_h28 = os.path.join(TMP, "28결과.html")
g.write_html(_h28, _res28, _sc28, _nt28, (P28, A28, S28X), set_name="28세트")
with open(_h28, encoding="utf-8") as f:
    _html28 = f.read()
check(28, "HTML 상단 배지 옆과 본문에 '확인 필요 N건' 섹션이 나온다",
      "확인 필요 1건" in _html28
      and '<div class="card review-card">' in _html28
      and '<div class="review-flag">' in _html28, "")
check(28, "HTML 에 무엇을 왜 판정 못 했는지와 점수 계산식이 적힌다",
      "vbaProject.bin" in _html28 and "점수 계산: 득점" in _html28
      and "채점된 배점" in _html28, "")
_resOK, _scOK, _ntOK, _ = g.run_grading(P28, A28, S28, g.normalize_key({}))
_hOK = os.path.join(TMP, "28확인필요없음.html")
g.write_html(_hOK, _resOK, _scOK, _ntOK, (P28, A28, S28), set_name="28세트")
_jOK = os.path.join(TMP, "28확인필요없음.json")
g.write_json(_jOK, _resOK, _scOK, _ntOK, (P28, A28, S28))
with open(_jOK, encoding="utf-8") as f:
    _dOK = json.load(f)
_htmlOK = open(_hOK, encoding="utf-8").read()
check(28, "확인 필요가 없으면 섹션도 표기도 없다 (정상 채점 방해 금지)",
      '<div class="card review-card">' not in _htmlOK
      and '<div class="review-flag">' not in _htmlOK
      and "확인 필요" not in _htmlOK.split("<style>")[0]
      and _dOK["review_count"] == 0 and _dOK["review_points"] == 0
      and _dOK["scoring"]["gradable_total"] == _dOK["scoring"]["alloc_total"],
      "")


# --- 28-5 실제 Excel 이 저장하는 마크업 형태 ---
# 아래 픽스처는 실제 Excel 이 내보낸 .xlsm 의 *마크업 형태*만 본떠 만든 것이다
# (내용은 이 테스트용으로 새로 씀). openpyxl 로 만든 합성 워크북은 이 형태를
# 쓰지 않아, 지금까지 이 경로는 한 번도 검증된 적이 없었다.
VML_EXCEL = """<xml xmlns:v="urn:schemas-microsoft-com:vml"
 xmlns:o="urn:schemas-microsoft-com:office:office"
 xmlns:x="urn:schemas-microsoft-com:office:excel">
 <o:shapelayout v:ext="edit">
  <o:idmap v:ext="edit" data="8"/>
 </o:shapelayout><v:shapetype id="_x0000_t201" coordsize="21600,21600" o:spt="201"
  path="m,l,21600r21600,l21600,xe">
  <v:stroke joinstyle="miter"/>
  <v:path shadowok="f" o:extrusionok="f" strokeok="f" fillok="f" o:connecttype="rect"/>
  <o:lock v:ext="edit" shapetype="t"/>
 </v:shapetype><v:shape id="_x0000_s8193" type="#_x0000_t201" style='position:absolute;
  margin-left:419.25pt;margin-top:49.5pt;width:108pt;height:33pt;z-index:1;
  mso-wrap-style:tight' o:button="t" fillcolor="buttonFace [67]"
  strokecolor="windowText [64]" o:insetmode="auto">
  <v:fill color2="buttonFace [67]" o:detectmouseclick="t"/>
  <o:lock v:ext="edit" rotation="t"/>
  <v:textbox style='mso-direction-alt:auto' o:singleclick="f">
   <div style='text-align:center'><font face="맑은 고딕" size="220" color="#000000">{text}<br
   />
      </font></div>
  </v:textbox>
  <x:ClientData ObjectType="Button">
   <x:Anchor>
    {anchor}</x:Anchor>
   <x:PrintObject>False</x:PrintObject>
   <x:AutoFill>False</x:AutoFill>
   <x:FmlaMacro>[0]!{macro}</x:FmlaMacro>
   <x:TextHAlign>Center</x:TextHAlign>
   <x:TextVAlign>Center</x:TextVAlign>
  </x:ClientData>
 </v:shape></xml>"""

_vml28 = VML_EXCEL.format(text="평균", macro="평균",
                          anchor="7, 0, 3, 0, 9, 0, 5, 0")
_shapes = re.findall(r"<v:shape\b.*?</v:shape>", _vml28, re.S)
check(28, "실제 Excel VML: <v:shapetype> 선행 블록을 단추로 잘못 세지 않는다",
      len(_shapes) == 1, f"{len(_shapes)}개")
check(28, "실제 Excel VML: 캡션 뒤 <br/> 가 줄바꿈으로 쪼개져 있어도 읽는다",
      g._vml_textbox_text(_shapes[0]) == "평균",
      repr(g._vml_textbox_text(_shapes[0])))
_anc = re.search(r"<x:Anchor>\s*([^<]*)</x:Anchor>", _shapes[0])
_nums = [n.strip() for n in _anc.group(1).split(",")]
check(28, "실제 Excel VML: <x:Anchor> 값이 줄바꿈+들여쓰기로 시작해도 읽는다",
      g._anchor_range(_nums[0], _nums[2], _nums[4], _nums[6]) == "H4:J6",
      g._anchor_range(_nums[0], _nums[2], _nums[4], _nums[6]))

# 시트 XML 의 <controls>: mc:AlternateContent + xdr: 접두사 + controlPr macro
CONTROLS_EXCEL = (
    '<controls><mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.'
    'org/markup-compatibility/2006"><mc:Choice Requires="x14">'
    '<control shapeId="8193" r:id="rIdCtrl" name="Button 1">'
    '<controlPr defaultSize="0" print="0" autoFill="0" autoPict="0" '
    'macro="[0]!평균"><anchor moveWithCells="1" sizeWithCells="1">'
    "<from><xdr:col>7</xdr:col><xdr:colOff>0</xdr:colOff>"
    "<xdr:row>3</xdr:row><xdr:rowOff>0</xdr:rowOff></from>"
    "<to><xdr:col>9</xdr:col><xdr:colOff>0</xdr:colOff>"
    "<xdr:row>5</xdr:row><xdr:rowOff>0</xdr:rowOff></to>"
    "</anchor></controlPr></control></mc:Choice></mc:AlternateContent>"
    "</controls>")


def inject_excel_button(src, out, with_vml=True, with_controls=True,
                        rel_target="../drawings/vmlDrawing1.vml",
                        text="평균", macro="평균",
                        sheet_part="xl/worksheets/sheet1.xml"):
    """실제 Excel 이 쓰는 형태로 양식 컨트롤 단추를 주입한다."""
    import zipfile
    with zipfile.ZipFile(src) as zin:
        blobs = {n: zin.read(n) for n in zin.namelist()}
    srels = re.sub(r"worksheets/(sheet\d+\.xml)$",
                   r"worksheets/_rels/\1.rels", sheet_part)
    rels = ['<Relationship Id="rIdCtrl" Type="http://schemas.openxmlformats.'
            'org/officeDocument/2006/relationships/ctrlProp" '
            'Target="../ctrlProps/ctrlProp1.xml"/>']
    if with_vml:
        rels.append('<Relationship Id="rIdVml" Type="http://schemas.'
                    'openxmlformats.org/officeDocument/2006/relationships/'
                    f'vmlDrawing" Target="{rel_target}"/>')
        blobs["xl/drawings/vmlDrawing1.vml"] = VML_EXCEL.format(
            text=text, macro=macro,
            anchor="7, 0, 3, 0, 9, 0, 5, 0").encode("utf-8")
    blobs["xl/ctrlProps/ctrlProp1.xml"] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<formControlPr xmlns="http://schemas.microsoft.com/office/'
        'spreadsheetml/2009/9/main" objectType="Button" lockText="1"/>'
    ).encode("utf-8")
    blobs[srels] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/'
        '2006/relationships">' + "".join(rels)
        + "</Relationships>").encode("utf-8")
    sx = blobs[sheet_part].decode("utf-8")
    if "xmlns:r=" not in sx.split(">", 1)[0]:
        sx = sx.replace("<worksheet ", '<worksheet xmlns:r="http://schemas.'
                        'openxmlformats.org/officeDocument/2006/'
                        'relationships" xmlns:xdr="http://schemas.'
                        'openxmlformats.org/drawingml/2006/'
                        'spreadsheetDrawing" ', 1)
    tail = ('<legacyDrawing r:id="rIdVml"/>' if with_vml else "") \
        + (CONTROLS_EXCEL if with_controls else "")
    blobs[sheet_part] = sx.replace(
        "</worksheet>", tail + "</worksheet>").encode("utf-8")
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for n, b in blobs.items():
            zout.writestr(n, b)
    return out


_x1 = inject_excel_button(save(wb_mac28(True), "28x1.xlsx"),
                          os.path.join(TMP, "28단추_실제Excel.xlsx"))
_b1 = g.Book(_x1, "테스트")
_fc1 = _b1.form_controls("매크로작업")
check(28, "실제 Excel 형태(VML+controls, ../ 상대 경로 rels)에서 단추를 "
          "한 개로 읽는다",
      not g.is_unknown(_fc1) and len(_fc1) == 1
      and _fc1[0]["text"] == "평균" and _fc1[0]["anchor"] == "H4:J6"
      and _fc1[0]["macro"] == "평균", str(_fc1))

# VML 없이 시트 <controls> 만 있는 경우도 읽어야 한다 (예전엔 앵커 정규식이
# xdr: 접두사를 못 넘고, 매크로도 controlPr 이 아닌 ctrlProps 에서만 찾아
# 앵커 '?' + 매크로 '' 로 버려졌다)
_x2 = inject_excel_button(save(wb_mac28(True), "28x2.xlsx"),
                          os.path.join(TMP, "28단추_controls만.xlsx"),
                          with_vml=False)
_fc2 = g.Book(_x2, "테스트").form_controls("매크로작업")
check(28, "시트 <controls> 만 있어도 xdr: 접두사 앵커와 controlPr macro 를 읽는다",
      not g.is_unknown(_fc2) and len(_fc2) == 1
      and _fc2[0]["anchor"] == "H4:J6" and _fc2[0]["macro"] == "평균",
      str(_fc2))

# 실제 Excel 의 drawing1.xml: mc:AlternateContent 로 감싼 hidden 단추 그림자
DRAWING_EXCEL = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/'
    'spreadsheetDrawing" xmlns:a="http://schemas.openxmlformats.org/'
    'drawingml/2006/main">'
    '<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/'
    'markup-compatibility/2006"><mc:Choice xmlns:a14="http://schemas.'
    'microsoft.com/office/drawing/2010/main" Requires="a14">'
    "<xdr:twoCellAnchor><xdr:from><xdr:col>7</xdr:col><xdr:colOff>0"
    "</xdr:colOff><xdr:row>3</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>"
    "<xdr:to><xdr:col>9</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>5"
    "</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>"
    '<xdr:sp macro="" textlink=""><xdr:nvSpPr>'
    '<xdr:cNvPr id="8193" name="Button 1" hidden="1"/><xdr:cNvSpPr/>'
    '</xdr:nvSpPr><xdr:spPr bwMode="auto"><a:prstGeom prst="rect">'
    "<a:avLst/></a:prstGeom></xdr:spPr><xdr:txBody><a:bodyPr/><a:lstStyle/>"
    "<a:p><a:r><a:rPr lang=\"ko-KR\"/><a:t>평균</a:t></a:r></a:p>"
    "<a:p><a:endParaRPr lang=\"ko-KR\"/></a:p></xdr:txBody></xdr:sp>"
    '<xdr:clientData fPrintsWithSheet="0"/></xdr:twoCellAnchor>'
    "</mc:Choice><mc:Fallback/></mc:AlternateContent>"
    "<xdr:twoCellAnchor><xdr:from><xdr:col>7</xdr:col><xdr:colOff>0"
    "</xdr:colOff><xdr:row>6</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>"
    "<xdr:to><xdr:col>9</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>8"
    "</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>"
    '<xdr:sp macro="[0]!서식" textlink=""><xdr:nvSpPr>'
    '<xdr:cNvPr id="2" name="다이아몬드 1"/><xdr:cNvSpPr/></xdr:nvSpPr>'
    '<xdr:spPr><a:prstGeom prst="diamond"><a:avLst/></a:prstGeom></xdr:spPr>'
    '<xdr:txBody><a:bodyPr/><a:lstStyle/><a:p><a:pPr algn="l"/><a:r>'
    '<a:rPr lang="ko-KR" altLang="en-US" sz="1100"/><a:t>서식</a:t></a:r>'
    "</a:p></xdr:txBody></xdr:sp><xdr:clientData/></xdr:twoCellAnchor>"
    "</xdr:wsDr>")


def inject_excel_drawing(src, out, sheet_part="xl/worksheets/sheet1.xml"):
    import zipfile
    with zipfile.ZipFile(src) as zin:
        blobs = {n: zin.read(n) for n in zin.namelist()}
    srels = re.sub(r"worksheets/(sheet\d+\.xml)$",
                   r"worksheets/_rels/\1.rels", sheet_part)
    cur = blobs.get(srels, (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/'
        '2006/relationships"></Relationships>').encode("utf-8")).decode("utf-8")
    blobs[srels] = cur.replace(
        "</Relationships>",
        '<Relationship Id="rIdDraw" Type="http://schemas.openxmlformats.org/'
        'officeDocument/2006/relationships/drawing" '
        'Target="../drawings/drawing1.xml"/></Relationships>'
    ).encode("utf-8")
    blobs["xl/drawings/drawing1.xml"] = DRAWING_EXCEL.encode("utf-8")
    sx = blobs[sheet_part].decode("utf-8")
    if "xmlns:r=" not in sx.split(">", 1)[0]:
        sx = sx.replace("<worksheet ", '<worksheet xmlns:r="http://schemas.'
                        'openxmlformats.org/officeDocument/2006/'
                        'relationships" ', 1)
    blobs[sheet_part] = sx.replace(
        "</worksheet>", '<drawing r:id="rIdDraw"/></worksheet>'
    ).encode("utf-8")
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for n, b in blobs.items():
            zout.writestr(n, b)
    return out


_x3 = inject_excel_drawing(save(wb_mac28(True), "28x3.xlsx"),
                           os.path.join(TMP, "28도형_실제Excel.xlsx"))
_sh3 = g.Book(_x3, "테스트").drawing_shapes("매크로작업")
check(28, "실제 Excel drawing: mc:AlternateContent 안의 앵커도 빠짐없이 읽는다",
      not g.is_unknown(_sh3) and len(_sh3) == 2, str(len(_sh3)))
check(28, "실제 Excel drawing: 단추 그림자(hidden)와 실제 도형을 구분한다",
      _sh3[0]["hidden"] and not _sh3[1]["hidden"]
      and _sh3[1]["geom"] == "diamond" and _sh3[1]["text"] == "서식"
      and _sh3[1]["macro"] == "서식", str(_sh3[1]))
_mo3 = g.macro_objects(g.Book(_x3, "테스트"), "매크로작업")
check(28, "매크로 개체 목록에서 hidden 그림자는 빼고 도형만 센다",
      len(_mo3) == 1 and _mo3[0]["kind"] == "도형", str(_mo3))

# 실제 Excel 이 저장하는 열 너비(문자 수 + 좌우 여백 5픽셀)
check(28, "실제 Excel 저장 열 너비를 대화상자 문자 수로 되돌린다 "
          "(맑은 고딕 11pt, MDW 8px)",
      all(g.col_width_is(stored, want) for stored, want in
          ((11.625, 11), (12.625, 12), (9.125, 8.5), (10.375, 10),
           (13.625, 13), (23.25, 23))),
      str([(s, g.col_width_chars(s)) for s in (11.625, 12.625, 23.25)]))
check(28, "Calibri 11pt(MDW 7px) 저장값도 같은 문자 수로 읽는다",
      g.col_width_is(12.7109375, 12) and g.col_width_is(8.7109375, 8), "")

# --- 28-6 남은 무음 실패 경로 ---
class _BoomBorderWS:
    """테두리를 읽는 순간 터지는 가짜 시트 (edge 맵 경로)."""

    merged_cells = type("R", (), {"ranges": []})()

    def cell(self, r, c):
        raise RuntimeError("테두리 파트 손상")


check(28, "sheet_edge_map: 테두리를 못 읽으면 '선 없음'이 아니라 Unknown "
          "(예전엔 None 으로 뭉개 테두리 감점으로 흘렀다)",
      g.is_unknown(g.sheet_edge_map(_BoomBorderWS(), 3, 3)),
      g.sheet_edge_map(_BoomBorderWS(), 3, 3).why)


class _BoomDxf:
    @property
    def font(self):
        raise RuntimeError("dxf 손상")


check(28, "_dxf_sig: 조건부 서식의 서식을 못 읽으면 Unknown",
      g.is_unknown(g._dxf_sig(_BoomDxf())), g._dxf_sig(_BoomDxf()).why)

_bad_excl = g.apply_key_exclusions(
    [], {}, [{"range": "이건범위가아님!!", "kinds": ["fill"]}])
check(28, "채점 제외 범위를 해석 못 하면 조용히 넘기지 않고 안내를 남긴다 "
          "(지시가 사라져 그대로 감점되던 경로)",
      any("해석하지 못해" in n for n in _bad_excl[2]), str(_bad_excl[2])[:90])

# 실제 Excel 이 스스로 만드는 예약 이름은 '정의된 이름' 지시가 아니다
check(28, "자동 필터·최신 함수가 만드는 예약 이름(_xlnm./_xleta.)은 "
          "정의된 이름 채점에서 제외한다",
      g._is_reserved_name("_xlnm._FilterDatabase")
      and g._is_reserved_name("_xleta.MINUTE")
      and not g._is_reserved_name("거래지점"), "")


# ---------------------------------------------------------------------------
# 29. v2.2.0 면책 문구 + 자가 채점 (채점기는 1차 의견, 최종 판정은 사용자)
#     이 채점기는 이의 신고 41건 중 25건(61%)이 오탐이었다. 배포판에는 그걸
#     전달할 곳이 없으니, 리포트가 (a) 틀릴 수 있다고 먼저 밝히고 (b) 쓰는
#     사람이 판정을 직접 뒤집을 수 있어야 한다.
#     '확인 필요'(v2.0.8 보류 = 채점기가 판단 못 함)와 자가 채점(= 사용자가
#     판정을 뒤집음)은 개념이 다르므로 리포트에서 따로 적어야 한다.
# ---------------------------------------------------------------------------
print("29. v2.2.0 면책 문구 · 자가 채점(채점기 1차 / 사용자 최종)")

import contextlib as _ctx29
import io as _io29
import re as _re29
import shutil as _sh29
import subprocess as _sp29


def _mk29(name, alloc, earned, cards=(), holds=()):
    r = g.SheetResult(name, alloc)
    r.earned = earned
    for _lbl, _lost in cards:
        g.add_card(r, _lbl, _lost, "value",
                   cells=[{"coord": "B2", "got": "1", "expected": "2"}])
    for _lbl, _why, _pts in holds:
        g.add_hold(r, _lbl, _why, _pts)
    return r


# 배점 100 · 득점 82 · 보류 6 -> 채점된 배점 94, 1차 점수 round(82/94*100)=87
_res29 = [
    _mk29("기본작업-1", 5, 5),
    _mk29("기본작업-2", 10, 6, cards=[("표시 형식", 4)]),
    _mk29("기본작업-3", 5, 5),
    _mk29("계산작업", 40, 32, cards=[("계산 문제 2", 8), ("설명 전용", 0)]),
    _mk29("분석작업-1", 10, 10),
    _mk29("분석작업-2", 10, 10),
    _mk29("매크로작업", 10, 4,
          holds=[("매크로 단추·도형", "vbaProject.bin 을 읽지 못했습니다", 6)]),
    _mk29("차트작업", 10, 10),
]
_alloc29 = sum(r.alloc for r in _res29)
_earn29 = sum(r.earned for r in _res29)
_held29 = g.hold_total(_res29)
_grad29 = _alloc29 - _held29
_sc29 = round(_earn29 / _grad29 * 100)
_P29 = os.path.join(TMP, "29세트_문제.xlsx")
_A29 = os.path.join(TMP, "29세트_정답.xlsx")
_S29 = os.path.join(TMP, "풀이_29세트_20260915_1226.xlsm")
_S29B = os.path.join(TMP, "풀이_29세트_20260916_0930.xlsm")
_H29 = os.path.join(TMP, "29리포트.html")
_H29B = os.path.join(TMP, "29리포트B.html")
_J29 = os.path.join(TMP, "29결과.json")
g.write_html(_H29, _res29, _sc29, ["안내 한 줄"], (_P29, _A29, _S29))
g.write_html(_H29B, _res29, _sc29, [], (_P29, _A29, _S29B))
g.write_json(_J29, _res29, _sc29, ["안내 한 줄"], (_P29, _A29, _S29))
with open(_H29, encoding="utf-8") as f:
    _html29 = f.read()
with open(_H29B, encoding="utf-8") as f:
    _html29b = f.read()
with open(_J29, encoding="utf-8") as f:
    _d29 = json.load(f)

# --- (a) 면책 문구 -------------------------------------------------------
_DISC = "채점기에 오류가 있을 수 있습니다. 참고용으로 봐 주세요."
check(29, "면책 문구 상수가 요구한 문장 그대로다",
      g.SELF_GRADE_DISCLAIMER == _DISC, g.SELF_GRADE_DISCLAIMER)
check(29, "리포트 상단(점수 바로 옆)에 면책 문구가 눈에 띄게 들어간다",
      'class="disclaimer"' in _html29
      and _DISC in _html29
      and _html29.index(_DISC) < _html29.index("영역별 점수"),
      "")
check(29, "면책 문구에 '실제 시험 채점과 다를 수 있다 · 직접 바꿀 수 있다'는 "
          "맥락 한 줄이 붙는다",
      "실제 시험 채점과 다를 수 있습니다" in _html29
      and "직접 판정을 바꿀 수 있고" in _html29)
check(29, "푸터에도 면책 문구를 남긴다", _html29.count(_DISC) >= 2,
      f"{_html29.count(_DISC)}회")
_buf29 = _io29.StringIO()
with _ctx29.redirect_stdout(_buf29):
    g.print_console(_res29, _sc29, [], (_P29, _A29, _S29))
check(29, "콘솔 출력에도 면책 문구 한 줄", _DISC in _buf29.getvalue())

# --- (b) 자가 채점 UI · 재계산 JS ---------------------------------------
_cards29 = [c for r in _res29 for c in r.wrong]
_scored29 = [c for c in _cards29 if c["lost"]]
check(29, "감점 있는 오답 카드마다 [내가 보기엔 맞음] 토글이 붙는다",
      _html29.count('class="sg-toggle"') == len(_scored29) == 2
      and _html29.count("data-sg-card=") == len(_scored29)
      and "내가 보기엔 맞음" in _html29,
      f"toggle={_html29.count('class=' + chr(34) + 'sg-toggle' + chr(34))}")
check(29, "감점 0점인 설명용 카드에는 토글을 달지 않는다 (눌러도 점수가 "
          "안 바뀌어 헷갈린다)",
      "되돌릴 감점이 없습니다" in _html29
      and _html29.count('class="wc"') == len(_cards29) == 3)
check(29, "카드마다 채점기의 1차 판정을 적어 둔다",
      "채점기 1차 판정: 틀림" in _html29 and "(-8점)" in _html29)
check(29, "점수를 '채점기 N점 (1차)' 와 '내가 매긴 점수 M점 (최종)' 두 값으로 "
          "나란히 적는다",
      'id="sg-auto"' in _html29 and 'id="sg-mine"' in _html29
      and "채점기 <strong" in _html29
      and "내가 매긴 점수 " in _html29
      and ">1차<" in _html29 and ">최종<" in _html29
      and f'id="sg-auto">{_sc29}<' in _html29
      and f'id="sg-mine">{_sc29}<' in _html29,
      f"1차={_sc29}")
check(29, "되돌리기(전체 초기화)와 [점수 복사] 경로가 있다",
      'id="sg-reset"' in _html29 and "전체 되돌리기" in _html29
      and 'id="sg-copy"' in _html29 and "점수 복사" in _html29)
check(29, "토글을 누르면 점수를 다시 계산하는 JS 가 들어 있다",
      "<script>" in _html29 and "KOCO_SELF_GRADE_DATA" in _html29
      and "window.kocoSelfGrade" in _html29
      and "roundHalfEven" in _html29
      and "addEventListener" in _html29)
check(29, "자가 채점 JS 는 저장소를 못 쓰는 환경에서도 죽지 않게 감싸 둔다",
      _html29.count("catch (e)") >= 6
      and "localStorage" in _html29)

# --- (c) 리포트별 localStorage 키 ----------------------------------------
_k29 = g.self_grade_key(_S29)
_k29b = g.self_grade_key(_S29B)
check(29, "저장 키는 리포트(응시)마다 갈린다", _k29 != _k29b
      and _k29 == "kocoSelfGrade:20260915122600"
      and _k29b == "kocoSelfGrade:20260916093000", f"{_k29} / {_k29b}")
check(29, "응시 시각이 없는 파일명도 리포트별로 안정된 키를 만든다 "
          "(채점 시각을 쓰면 다시 채점할 때마다 선택이 날아간다)",
      g.self_grade_key("내풀이.xlsx") == g.self_grade_key("내풀이.xlsx")
      and g.self_grade_key("내풀이.xlsx") != g.self_grade_key("다른풀이.xlsx")
      and g.self_grade_key("내풀이.xlsx").startswith("kocoSelfGrade:"),
      g.self_grade_key("내풀이.xlsx"))
check(29, "리포트 안에는 자기 키만 들어간다 (다른 회차 선택이 섞이지 않는다)",
      _k29 in _html29 and _k29b not in _html29
      and _k29b in _html29b and _k29 not in _html29b)
check(29, "제거한 이의제기 기능의 kocoAppeal 키 관례는 쓰지 않는다",
      "kocoAppeal" not in _html29
      and g.SELF_GRADE_KEY_PREFIX == "kocoSelfGrade:")

# --- (d) '확인 필요'(보류)와 자가 채점을 구분해 적는다 -------------------
check(29, "확인 필요 항목은 자가 채점 카드와 다른 칸(data-sg-hold)에 적는다",
      _html29.count("data-sg-hold=") == 1
      and _html29.count("data-sg-card=") == 2
      and '"card review-card"' in _html29)
check(29, "두 개념의 차이를 리포트가 말로 설명한다 "
          "(보류=채점기가 판단 못 함 / 자가 채점=사용자가 판정을 뒤집음)",
      "자가 채점과는 다른 칸" in _html29
      and "채점기가 아예 판단하지 못한" in _html29
      and "채점기가 '틀림'으로 판정한 것을 사용자가" in _html29)
check(29, "확인 필요 항목에는 [맞음]/[틀림]/[채점 제외로 두기]를 둔다 "
          "(고르면 그 배점이 채점 대상으로 돌아온다)",
      'data-sg-hold-val="right"' in _html29
      and 'data-sg-hold-val="wrong"' in _html29
      and "채점 제외로 두기" in _html29
      and "채점기가 판단하지 못한 항목" in _html29)
check(29, "확인 필요의 기본값은 그대로 '점수에서 뺌' (v2.0.8 규약 유지)",
      "확인 필요 1건" in _html29 and "점수에서 뺌" in _html29
      and f"× 100 = <strong>{_sc29}점</strong>" in _html29)

# --- (e) 재계산 산식: 리포트 payload 가 파이썬 채점 결과와 맞는가 --------
_pay29 = json.loads(_re29.search(
    r"window\.KOCO_SELF_GRADE_DATA=(\{.*?\});\n", _html29, _re29.S).group(1)
    .replace("\\u003c", "<").replace("\\u003e", ">").replace("\\u0026", "&"))
check(29, "리포트 JS 에 넘기는 1차 점수·득점·채점된 배점이 파이썬 채점과 같다",
      _pay29["baseScore"] == _sc29 and _pay29["earned"] == _earn29
      and _pay29["gradable"] == _grad29 and _pay29["maxTotal"] == 100
      and _pay29["partial"] is False,
      json.dumps({k: _pay29[k] for k in ("baseScore", "earned", "gradable")}))


def _recalc29(flip_ids, hold_verdicts):
    """리포트 JS 와 같은 식으로 다시 매긴 점수 (파이썬 거울)."""
    earned, gradable = _pay29["earned"], _pay29["gradable"]
    for c in _pay29["cards"]:
        if c["id"] in flip_ids:
            earned += c["lost"]
    for h in _pay29["holds"]:
        v = hold_verdicts.get(h["id"])
        if v == "right":
            earned += h["points"]
            gradable += h["points"]
        elif v == "wrong":
            gradable += h["points"]
    return min(100, max(0, round(earned / gradable * 100)))


_ids29 = [c["id"] for c in _pay29["cards"]]
_hid29 = [h["id"] for h in _pay29["holds"]]
check(29, "오답을 하나 뒤집으면 그 감점만큼 점수가 오른다",
      _recalc29({_ids29[1]}, {}) == round((82 + 8) / 94 * 100) == 96,
      str(_recalc29({_ids29[1]}, {})))
check(29, "오답을 전부 뒤집으면 감점이 전부 되돌아온다",
      _recalc29(set(_ids29), {}) == 100, str(_recalc29(set(_ids29), {})))
check(29, "확인 필요를 '맞음'으로 판정하면 그 배점이 분모로 돌아오고 득점에 "
          "더해진다", _recalc29(set(), {_hid29[0]: "right"})
      == round(88 / 100 * 100) == 88,
      str(_recalc29(set(), {_hid29[0]: "right"})))
check(29, "확인 필요를 '틀림'으로 판정하면 분모만 돌아와 점수가 내려간다",
      _recalc29(set(), {_hid29[0]: "wrong"}) == round(82 / 100 * 100) == 82,
      str(_recalc29(set(), {_hid29[0]: "wrong"})))
check(29, "아무것도 안 바꾸면 채점기의 1차 점수 그대로",
      _recalc29(set(), {}) == _sc29)

# --- (f) --json: 기존 키 보존 + self_grading 추가 ------------------------
check(29, "JSON 기존 키는 그대로 (시험장.py 가 읽는다)",
      all(k in _d29 for k in ("total", "pass_line", "passed", "mode",
                              "graded_sheets", "max_total", "generated",
                              "files", "sheets", "notes", "wrong_items",
                              "review_count", "review_points", "review_items",
                              "scoring"))
      and _d29["total"] == _sc29
      and len(_d29["wrong_items"]) == len(_cards29)
      and _d29["review_count"] == 1, "")
check(29, "JSON 에 self_grading 을 추가한다 (추가만 — 기존 키 의미는 그대로)",
      "self_grading" in _d29
      and _d29["self_grading"]["storageKey"] == _k29
      and _d29["self_grading"]["baseScore"] == _sc29
      and _d29["self_grading"]["disclaimer"] == _DISC
      and "1차 의견" in _d29["self_grading"]["method"],
      _d29["self_grading"]["storageKey"])
check(29, "JSON 의 자가 채점 항목 id 가 리포트 HTML 의 것과 같다 "
          "(둘을 맞대 볼 수 있어야 한다)",
      [c["id"] for c in _d29["self_grading"]["cards"]] == _ids29
      and [h["id"] for h in _d29["self_grading"]["holds"]] == _hid29
      and all(f'data-sg-card="{i}"' in _html29 for i in _ids29 if i in
              {c["id"] for c in _pay29["cards"] if c["lost"]}),
      "")

# --- (g) 만점 리포트 / v2.1.0 제약 유지 ----------------------------------
_resP29 = [_mk29("기본작업-1", 5, 5), _mk29("계산작업", 95, 95)]
_HP29 = os.path.join(TMP, "29만점.html")
g.write_html(_HP29, _resP29, 100, [], (_P29, _A29, _S29))
with open(_HP29, encoding="utf-8") as f:
    _htmlP29 = f.read()
check(29, "만점 리포트에도 면책 문구와 점수 패널은 있고 토글은 없다",
      _DISC in _htmlP29 and 'id="sg-mine"' in _htmlP29
      and "틀린 항목이 없습니다" in _htmlP29
      and 'class="sg-toggle"' not in _htmlP29
      and "data-sg-card=" not in _htmlP29
      and "data-sg-hold=" not in _htmlP29)
check(29, "v2.1.0 제약 유지: 이의제기 흔적·execCommand 없음, <body> 그대로",
      "이의제기" not in _html29 and "appeal" not in _html29
      and "execCommand" not in _html29 and "<body>" in _html29
      and "이의제기" not in _htmlP29 and "appeal" not in _htmlP29)

# --- (h) JS 문법 검사 (리포트는 사용자 PC 브라우저에서 도는 코드다) ------
_js29 = _re29.search(r"<script>(.*?)</script>", _html29, _re29.S).group(1)
check(29, "리포트 안의 <script> 가 따옴표·이스케이프로 깨지지 않았다",
      _js29.count("</script") == 0 and "\\u003c" not in _js29.split(";\n")[1:2]
      and _js29.lstrip().startswith("window.KOCO_SELF_GRADE_DATA=")
      and _js29.rstrip().endswith("})();"), _js29[:60])
_node29 = _sh29.which("node")
if _node29:
    _jsf29 = os.path.join(TMP, "29리포트.js")
    with open(_jsf29, "w", encoding="utf-8") as f:
        f.write(_js29)
    _cp29 = _sp29.run([_node29, "--check", _jsf29], capture_output=True,
                      text=True)
    check(29, "node --check 로 리포트 JS 문법 검사 통과",
          _cp29.returncode == 0, (_cp29.stderr or "")[-160:])
else:
    print("  [29] node 가 없어 JS 문법 검사는 건너뜁니다 (문자열 검사만 수행)")

print()
print("전체 회귀 스위트 통과 (오탐 7계열 + 감사 케이스 + v1.9.0 정밀화 "
      "+ v2.1.0 이의제기 제거 + v2.0.1 이의제기 반영 + v2.0.2 이의제기 반영 "
      "+ v2.0.3 이의제기 반영 + v2.0.4 이의제기 반영 "
      "+ v2.0.5 이의제기 반영 + v2.0.6 이의제기 반영 "
      "+ v2.0.8 UNKNOWN 보류·확인 필요 표기·실제 Excel 마크업 "
      "+ v2.2.0 면책 문구·자가 채점(채점기 1차 / 사용자 최종))")
